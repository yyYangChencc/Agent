from __future__ import annotations

import argparse
import csv
import html
import json
import math
import random
import sys
from collections import Counter, defaultdict
from dataclasses import dataclass
from pathlib import Path
from statistics import mean, median
from typing import Any, Iterable

from persona.history_csv import configure_csv_field_size_limit
from persona.opinion.scale import (
    VOTING_ROLE_OPPOSE,
    VOTING_ROLE_SUPPORT,
    VOTING_ROLE_UNKNOWN,
    VOTING_STANCE_INVALID,
    classify_voting_stance,
)


def _configure_csv_field_size_limit() -> int:
    """兼容旧调用入口，并复用统一的 CSV 字段上限配置。"""

    return configure_csv_field_size_limit()


_configure_csv_field_size_limit()


HISTORY_DIR = Path(__file__).resolve().parent / "history"
REQUIRED_HISTORY_COLUMNS = ("tick", "opinion")
OPENPYXL_SUFFIXES = {".xlsx", ".xlsm", ".xltx", ".xltm"}
XLRD_SUFFIXES = {".xls"}
CSV_SUFFIXES = {".csv"}
TABLE_SUFFIXES = OPENPYXL_SUFFIXES | XLRD_SUFFIXES | CSV_SUFFIXES
INVALID_FILENAME_CHARS = '<>:"/\\|?*'
GENERATED_ANALYSIS_FILES = {
    "polarization_metrics.csv",
    "polarization_agent_shift.csv",
    "opinion_voting_polarization_metrics.csv",
    "opinion_voting_agent_metrics.csv",
    "opinion_voting_skipped_records.csv",
    "opinion_distribution.csv",
}


@dataclass(frozen=True)
class AnalysisOptions:
    data_length: int | None = None
    opinion_baseline_window_size: int = 10
    opinion_final_window_size: int = 10
    voting_baseline_window_size: int = 1
    voting_final_window_size: int = 1
    support_threshold: float = 0.35
    oppose_threshold: float = -0.35
    neutral_threshold: float = 0.10
    extreme_threshold: float = 0.75
    min_side_share: float = 0.20
    min_pairwise_delta: float = 0.05
    min_abs_delta: float = 0.05
    alpha: float = 0.05
    bootstrap_samples: int = 2000
    seed: int = 42
    allow_incomplete_ticks: bool = False
    min_voting_known_share: float = 0.70
    min_voting_decisiveness: float = 0.60
    min_voting_polarization_delta: float = 0.05


@dataclass(frozen=True)
class AnalysisResult:
    run_dir: Path
    output_dir: Path
    summary_path: Path
    polarization_report_path: Path | None
    generated_paths: list[Path]


@dataclass(frozen=True)
class TickMetric:
    tick: int
    n_agents: int
    mean_opinion: float
    median_opinion: float
    std_population: float
    mean_abs_opinion: float
    median_abs_opinion: float
    pairwise_mean_abs_distance: float
    min_opinion: float
    max_opinion: float
    range_opinion: float
    support_count: int
    oppose_count: int
    neutral_count: int
    extreme_count: int
    support_share: float
    oppose_share: float
    neutral_share: float
    extreme_share: float
    two_side_share: float
    camp_gap: float | None
    polarization_index: float


@dataclass(frozen=True)
class AgentShift:
    agent_id: str
    baseline_mean_opinion: float
    final_mean_opinion: float
    delta_opinion: float
    baseline_mean_abs_opinion: float
    final_mean_abs_opinion: float
    delta_abs_opinion: float


@dataclass(frozen=True)
class VotingAgentMetric:
    tick: int
    agent_id: str
    requested_votes: int
    successful_votes: int
    failed_votes: int
    skipped_reason: str
    support_share: float
    oppose_share: float
    unknown_share: float
    known_share: float
    decisiveness: float
    stance: str
    stance_valid: int
    stance_agreement: float
    stance_direction_margin: float
    stance_success_rate: float
    previous_valid_tick: int
    previous_valid_stance: str
    stance_changed: int


@dataclass(frozen=True)
class VotingSkippedRecord:
    tick: int | None
    agent_id: str
    reason: str
    requested_votes: int | None
    successful_votes: int | None
    failed_votes: int | None


@dataclass(frozen=True)
class VotingTickMetric:
    tick: int
    n_agents: int
    valid_agents: int
    requested_votes: int
    successful_votes: int
    failed_votes: int
    success_rate: float
    support_share: float
    oppose_share: float
    unknown_share: float
    known_share: float
    decisiveness: float
    camp_balance: float
    polarization_index: float
    stance_valid_agents: int
    stance_invalid_agents: int
    stance_support_count: int
    stance_oppose_count: int
    stance_unknown_count: int
    stance_support_share: float
    stance_oppose_share: float
    stance_unknown_share: float
    stance_invalid_share: float
    stance_comparable_agents: int
    stance_changed_agents: int
    stance_unchanged_agents: int
    stance_changed_share: float
    stance_changed_to_support: int
    stance_changed_to_oppose: int
    stance_changed_to_unknown: int


@dataclass(frozen=True)
class PlotSeries:
    label: str
    x_values: list[Any]
    y_values: list[float]


def parse_args(argv: list[str]) -> argparse.Namespace:
    parser = argparse.ArgumentParser(
        description="统一分析 history 数据，生成实验摘要、曲线图、极化统计，也可按指定列绘图。",
        epilog=(
            "示例：python backend/analyze_history.py backend/history/20260628_181756\n"
            "示例：python backend/analyze_history.py backend/history/20260628_181756 --plot-columns opinion --x-column tick"
        ),
    )
    parser.add_argument("history_dir", nargs="?", default=str(HISTORY_DIR), help="history 运行目录，或 history 根目录。")
    parser.add_argument("--output-dir", help="输出目录；默认写入被分析的运行目录。")
    parser.add_argument("--data-length", type=int, help="只分析或绘制开头 N 条数据。")
    parser.add_argument(
        "--window-size",
        type=int,
        help="兼容参数：同时设置连续 opinion 的基线和末端窗口长度。",
    )
    parser.add_argument(
        "--opinion-baseline-window-size",
        type=int,
        help="连续 opinion 基线窗口的 tick 数，默认 10。",
    )
    parser.add_argument(
        "--opinion-final-window-size",
        type=int,
        help="连续 opinion 末端窗口的 tick 数，默认 10。",
    )
    parser.add_argument(
        "--voting-baseline-window-size",
        type=int,
        default=1,
        help="llm_voting 基线窗口的评测轮数，默认 1。",
    )
    parser.add_argument(
        "--voting-final-window-size",
        type=int,
        default=1,
        help="llm_voting 末端窗口的评测轮数，默认 1。",
    )
    parser.add_argument("--support-threshold", type=float, default=0.35, help="明显支持阈值，默认 0.35。")
    parser.add_argument("--oppose-threshold", type=float, default=-0.35, help="明显反对阈值，默认 -0.35。")
    parser.add_argument("--neutral-threshold", type=float, default=0.10, help="中立区间绝对值阈值，默认 0.10。")
    parser.add_argument("--extreme-threshold", type=float, default=0.75, help="极端立场绝对值阈值，默认 0.75。")
    parser.add_argument("--min-side-share", type=float, default=0.20, help="双边阵营最小占比，默认 0.20。")
    parser.add_argument("--min-pairwise-delta", type=float, default=0.05, help="平均成对距离最小增长，默认 0.05。")
    parser.add_argument("--min-abs-delta", type=float, default=0.05, help="平均绝对立场最小增长，默认 0.05。")
    parser.add_argument("--alpha", type=float, default=0.05, help="配对符号检验显著性水平，默认 0.05。")
    parser.add_argument("--bootstrap-samples", type=int, default=2000, help="配对均值差 bootstrap 次数，默认 2000。")
    parser.add_argument("--seed", type=int, default=42, help="bootstrap 随机种子，默认 42。")
    parser.add_argument("--allow-incomplete-ticks", action="store_true", help="允许部分智能体缺失的 tick 参与极化统计。")
    parser.add_argument("--min-voting-known-share", type=float, default=0.70, help="投票模式有效立场覆盖率下限，默认 0.70。")
    parser.add_argument("--min-voting-decisiveness", type=float, default=0.60, help="投票模式个体明确度下限，默认 0.60。")
    parser.add_argument("--min-voting-polarization-delta", type=float, default=0.05, help="投票极化指数最小增长，默认 0.05。")
    parser.add_argument("--plot-columns", nargs="+", help="按精确列名绘制所有 CSV/Excel 文件中的对应列。")
    parser.add_argument("--x-column", help="列绘图横轴列名，必须与表头完全一致；省略时使用数据行序号。")
    parser.add_argument("--sheet", help="列绘图只读取指定 Excel 工作表；CSV 不支持该参数。")
    parser.add_argument("--header-row", type=int, default=1, help="列绘图表头所在行号，默认第 1 行。")
    parser.add_argument("--format", default="svg", choices=("svg", "png", "pdf"), help="列绘图输出格式，默认 svg。")
    parser.add_argument("--no-recursive", action="store_true", help="列绘图只读取直属 CSV/Excel 文件。")
    parser.add_argument("--list-columns", action="store_true", help="只列出 CSV/Excel 表头，不生成统计。")
    parser.add_argument("--no-legend", action="store_true", help="列绘图不显示图例。")
    parser.add_argument("--legend-limit", type=int, default=30, help="列绘图曲线数量不超过该值时显示图例，默认 30。")
    args = parser.parse_args(normalize_argv(argv))
    validate_args(parser, args)
    return args


def normalize_argv(argv: list[str]) -> list[str]:
    normalized: list[str] = []
    index = 0
    while index < len(argv):
        item = argv[index]
        if item.startswith("--") and item[2:].isdigit():
            normalized.extend(["--data-length", item[2:]])
            if index + 1 < len(argv) and not argv[index + 1].startswith("-"):
                normalized.extend(["--x-column", argv[index + 1]])
                index += 2
                continue
        else:
            normalized.append(item)
        index += 1
    return normalized


def validate_args(parser: argparse.ArgumentParser, args: argparse.Namespace) -> None:
    if args.data_length is not None and args.data_length < 1:
        parser.error("--data-length 必须大于等于 1。")
    opinion_baseline_size = (
        args.opinion_baseline_window_size
        if args.opinion_baseline_window_size is not None
        else args.window_size if args.window_size is not None else 10
    )
    opinion_final_size = (
        args.opinion_final_window_size
        if args.opinion_final_window_size is not None
        else args.window_size if args.window_size is not None else 10
    )
    if args.window_size is not None and args.window_size < 1:
        parser.error("--window-size 必须大于等于 1。")
    if opinion_baseline_size < 1:
        parser.error("--opinion-baseline-window-size 必须大于等于 1。")
    if opinion_final_size < 1:
        parser.error("--opinion-final-window-size 必须大于等于 1。")
    if args.voting_baseline_window_size < 1:
        parser.error("--voting-baseline-window-size 必须大于等于 1。")
    if args.voting_final_window_size < 1:
        parser.error("--voting-final-window-size 必须大于等于 1。")
    if args.header_row < 1:
        parser.error("--header-row 必须大于等于 1。")
    if args.legend_limit < 0:
        parser.error("--legend-limit 必须大于等于 0。")
    if not 0 <= args.min_side_share <= 1:
        parser.error("--min-side-share 必须在 [0, 1] 内。")
    if not 0 < args.alpha < 1:
        parser.error("--alpha 必须在 (0, 1) 内。")
    if args.bootstrap_samples < 0:
        parser.error("--bootstrap-samples 必须大于等于 0。")
    if not 0 <= args.min_voting_known_share <= 1:
        parser.error("--min-voting-known-share 必须在 [0, 1] 内。")
    if not 0 <= args.min_voting_decisiveness <= 1:
        parser.error("--min-voting-decisiveness 必须在 [0, 1] 内。")
    if not 0 <= args.min_voting_polarization_delta <= 1:
        parser.error("--min-voting-polarization-delta 必须在 [0, 1] 内。")
    if args.support_threshold <= args.oppose_threshold:
        parser.error("--support-threshold 必须大于 --oppose-threshold。")


def main(argv: list[str] | None = None) -> int:
    args = parse_args(sys.argv[1:] if argv is None else argv)
    history_dir = Path(args.history_dir).expanduser()
    output_dir = Path(args.output_dir).expanduser() if args.output_dir else None

    if args.list_columns:
        table_paths = find_table_files(history_dir, recursive=not args.no_recursive)
        print_column_report(table_paths, args.sheet, args.header_row)
        return 0

    if args.plot_columns:
        paths = plot_history_columns(
            history_dir=history_dir,
            columns=args.plot_columns,
            output_dir=output_dir,
            x_column=args.x_column,
            sheet_name=args.sheet,
            header_row=args.header_row,
            image_format=args.format,
            data_length=args.data_length,
            recursive=not args.no_recursive,
            show_legend=not args.no_legend,
            legend_limit=args.legend_limit,
        )
        for path in paths:
            print(f"已生成：{path}")
        return 0

    result = analyze_history_run(
        history_dir,
        output_dir=output_dir,
        options=AnalysisOptions(
            data_length=args.data_length,
            opinion_baseline_window_size=(
                args.opinion_baseline_window_size
                if args.opinion_baseline_window_size is not None
                else args.window_size if args.window_size is not None else 10
            ),
            opinion_final_window_size=(
                args.opinion_final_window_size
                if args.opinion_final_window_size is not None
                else args.window_size if args.window_size is not None else 10
            ),
            voting_baseline_window_size=args.voting_baseline_window_size,
            voting_final_window_size=args.voting_final_window_size,
            support_threshold=args.support_threshold,
            oppose_threshold=args.oppose_threshold,
            neutral_threshold=args.neutral_threshold,
            extreme_threshold=args.extreme_threshold,
            min_side_share=args.min_side_share,
            min_pairwise_delta=args.min_pairwise_delta,
            min_abs_delta=args.min_abs_delta,
            alpha=args.alpha,
            bootstrap_samples=args.bootstrap_samples,
            seed=args.seed,
            allow_incomplete_ticks=args.allow_incomplete_ticks,
            min_voting_known_share=args.min_voting_known_share,
            min_voting_decisiveness=args.min_voting_decisiveness,
            min_voting_polarization_delta=args.min_voting_polarization_delta,
        ),
    )
    for path in result.generated_paths:
        print(f"已生成：{path}")
    return 0


def analyze_history_run(
    history_dir: str | Path,
    *,
    output_dir: str | Path | None = None,
    config_snapshot: dict | None = None,
    options: AnalysisOptions | None = None,
) -> AnalysisResult:
    """统一生成实验摘要、基础曲线和极化统计。"""

    opts = options or AnalysisOptions()
    run_dir = resolve_history_run_dir(history_dir)
    out_dir = Path(output_dir).expanduser() if output_dir is not None else run_dir
    out_dir.mkdir(parents=True, exist_ok=True)
    rows_by_agent = read_history_rows(run_dir)
    snapshot = resolve_config_snapshot(run_dir, config_snapshot)
    assessment_mode = resolve_opinion_assessment_mode(snapshot, rows_by_agent)

    summary_path = out_dir / "experiment_summary.md"
    summary_path.write_text(
        build_experiment_summary(rows_by_agent, snapshot),
        encoding="utf-8",
    )
    generated_paths = [summary_path]
    generated_paths.extend(
        write_standard_charts(
            out_dir,
            rows_by_agent,
            include_voting_chart=False,
        )
    )
    generated_paths.extend(
        write_opinion_analysis_outputs(
            run_dir=run_dir,
            output_dir=out_dir,
            options=opts,
        )
    )

    polarization_report_path = None
    if assessment_mode == "llm_voting":
        try:
            posthoc_voting_rows = read_posthoc_voting_rows(run_dir)
            voting_rows = posthoc_voting_rows or rows_by_agent
            generated_paths.append(
                write_opinion_voting_trend_svg(
                    out_dir / "opinion_voting_trends.svg",
                    voting_rows,
                )
            )
            polarization_paths = write_voting_polarization_outputs(
                rows_by_agent=voting_rows,
                output_dir=out_dir,
                options=opts,
                expected_agent_ids=set(rows_by_agent),
            )
            polarization_report_path = out_dir / "opinion_voting_polarization_report.md"
            generated_paths.extend(polarization_paths)
        except ValueError as exc:
            polarization_report_path = _write_polarization_error_report(
                out_dir / "opinion_voting_polarization_report.md",
                exc,
                title="LLM 投票极化统计报告",
            )
            generated_paths.append(polarization_report_path)
    else:
        try:
            polarization_paths = write_polarization_outputs(
                run_dir=run_dir,
                output_dir=out_dir,
                options=opts,
            )
            polarization_report_path = out_dir / "polarization_report.md"
            generated_paths.extend(polarization_paths)
        except ValueError as exc:
            polarization_report_path = _write_polarization_error_report(
                out_dir / "polarization_report.md",
                exc,
                title="舆论极化统计报告",
            )
            generated_paths.append(polarization_report_path)

        # 附加投票链独立失败时，不得覆盖已经生成的连续观念报告。
        try:
            posthoc_voting_rows = read_posthoc_voting_rows(run_dir)
            generated_paths.append(
                write_opinion_voting_trend_svg(
                    out_dir / "opinion_voting_trends.svg",
                    posthoc_voting_rows or rows_by_agent,
                )
            )
            if posthoc_voting_rows:
                generated_paths.extend(
                    write_voting_polarization_outputs(
                        rows_by_agent=posthoc_voting_rows,
                        output_dir=out_dir,
                        options=opts,
                        expected_agent_ids=set(rows_by_agent),
                    )
                )
        except ValueError as exc:
            generated_paths.append(
                _write_polarization_error_report(
                    out_dir / "opinion_voting_polarization_report.md",
                    exc,
                    title="LLM 投票极化统计报告",
                )
            )

    return AnalysisResult(
        run_dir=run_dir,
        output_dir=out_dir,
        summary_path=summary_path,
        polarization_report_path=polarization_report_path,
        generated_paths=generated_paths,
    )


def _write_polarization_error_report(path: Path, exc: ValueError, *, title: str) -> Path:
    """将单类极化分析错误写入自己的报告，避免覆盖其他分析结果。"""

    path.write_text(f"# {title}\n\n无法生成极化统计：{exc}\n", encoding="utf-8")
    return path


def resolve_config_snapshot(run_dir: Path, config_snapshot: dict | None) -> dict:
    """优先使用调用方快照，否则读取运行目录中的精确快照文件。"""

    if config_snapshot:
        return dict(config_snapshot)
    path = run_dir / "config_snapshot.json"
    if not path.exists():
        return {}
    try:
        payload = json.loads(path.read_text(encoding="utf-8-sig"))
    except (OSError, json.JSONDecodeError) as exc:
        raise ValueError(f"无法读取配置快照 {path}：{exc}") from exc
    if not isinstance(payload, dict):
        raise ValueError(f"配置快照根节点必须是对象：{path}")
    return payload


def read_posthoc_voting_rows(run_dir: Path) -> dict[str, list[dict]]:
    """把结束后投票 JSON 转成现有投票分析器使用的精确字段。"""

    path = run_dir / "opinion_voting_posthoc.json"
    if not path.exists():
        return {}
    try:
        payload = json.loads(path.read_text(encoding="utf-8-sig"))
    except (OSError, json.JSONDecodeError) as exc:
        raise ValueError(f"无法读取结束后投票结果 {path}：{exc}") from exc
    if not isinstance(payload, list):
        raise ValueError(f"结束后投票结果根节点必须是数组：{path}")
    rows_by_agent: dict[str, list[dict]] = defaultdict(list)
    for item in payload:
        if not isinstance(item, dict):
            raise ValueError(f"结束后投票结果包含非对象元素：{path}")
        agent_id = item.get("agent_id")
        if not isinstance(agent_id, str) or not agent_id:
            raise ValueError(f"结束后投票结果缺少 agent_id：{path}")
        rows_by_agent[agent_id].append({
            "opinion_assessment_topic": item.get("topic", ""),
            "opinion_voting_tick": item.get("tick", ""),
            "opinion_voting_choice_counts": json.dumps(item.get("choice_counts", {}), ensure_ascii=False),
            "opinion_voting_option_roles": json.dumps(item.get("option_roles", {}), ensure_ascii=False),
            "opinion_voting_requested_voters": item.get("requested_voters", ""),
            "opinion_voting_successful_votes": item.get("successful_votes", ""),
            "opinion_voting_failed_votes": item.get("failed_votes", ""),
            "opinion_voting_skipped_reason": item.get("skipped_reason", ""),
            "opinion_voting_stance": item.get("stance", ""),
            "opinion_voting_stance_valid": item.get("stance_valid", ""),
            "opinion_voting_options": json.dumps(item.get("options", []), ensure_ascii=False),
            "opinion_voting_window_start_tick": item.get("window_start_tick", ""),
            "opinion_voting_window_end_tick": item.get("window_end_tick", ""),
            "opinion_voting_speech_history": json.dumps(item.get("speech_history", []), ensure_ascii=False),
        })
    return dict(rows_by_agent)


def resolve_opinion_assessment_mode(config_snapshot: dict, rows_by_agent: dict[str, list[dict]]) -> str:
    """按配置快照或历史记录中的精确方法字段确定分析方式。"""

    mode = config_snapshot.get("opinion_assessment_mode")
    if mode in {"llm_voting", "llm_as_judge", "rule"}:
        return str(mode)
    recorded_modes = {
        str(row.get("opinion_assessment_method"))
        for rows in rows_by_agent.values()
        for row in rows
        if row.get("opinion_assessment_method") in {"llm_voting", "llm_as_judge", "rule_context_assessment"}
    }
    if "llm_voting" in recorded_modes:
        return "llm_voting"
    return "llm_as_judge"


def resolve_history_run_dir(history_dir: str | Path) -> Path:
    path = Path(history_dir).expanduser()
    if not path.exists():
        raise FileNotFoundError(f"history 路径不存在：{path}")
    if not path.is_dir():
        raise NotADirectoryError(f"history 路径不是文件夹：{path}")
    if history_csv_paths(path):
        return path

    run_dirs = sorted(
        (child for child in path.iterdir() if child.is_dir()),
        key=lambda child: child.stat().st_mtime,
        reverse=True,
    )
    for run_dir in run_dirs:
        if history_csv_paths(run_dir):
            return run_dir
    raise FileNotFoundError(f"没有找到包含 {REQUIRED_HISTORY_COLUMNS} 的智能体历史 CSV：{path}")


def read_history_rows(run_dir: Path) -> dict[str, list[dict]]:
    rows_by_agent = {
        path.stem: read_csv_rows(path)
        for path in history_csv_paths(run_dir)
    }
    if not rows_by_agent:
        raise ValueError(f"没有读取到包含 {REQUIRED_HISTORY_COLUMNS} 的智能体历史 CSV：{run_dir}")
    return rows_by_agent


def history_csv_paths(run_dir: Path) -> list[Path]:
    """返回真正的智能体历史 CSV；文件名可使用任意精确 agent_id。"""

    return [
        path
        for path in sorted(run_dir.glob("*.csv"))
        if is_history_agent_csv(path)
    ]


def is_history_agent_csv(path: Path) -> bool:
    """通过表头识别历史 CSV，避免把分析产物再次读入。"""

    if path.name in GENERATED_ANALYSIS_FILES or path.name.startswith("~$"):
        return False
    try:
        with path.open("r", newline="", encoding="utf-8-sig") as f:
            reader = csv.reader(f)
            header = next(reader, [])
    except (OSError, UnicodeDecodeError):
        return False
    return all(column in header for column in REQUIRED_HISTORY_COLUMNS)


def read_csv_rows(path: Path) -> list[dict]:
    with path.open("r", newline="", encoding="utf-8-sig") as f:
        return list(csv.DictReader(f))


def build_experiment_summary(rows_by_agent: dict[str, list[dict]], config_snapshot: dict) -> str:
    voting_mode = config_snapshot.get("opinion_assessment_mode") == "llm_voting"
    final_opinions = {}
    pressure_peaks = defaultdict(dict)
    mediator_peaks = defaultdict(float)
    behavior_counts = Counter()
    post_scores = []
    abnormal_ticks = []
    vote_totals = Counter()
    failed_vote_total = 0

    for agent_id, rows in rows_by_agent.items():
        if not rows:
            continue
        final = rows[-1]
        final_opinions[agent_id] = parse_float(final.get("opinion"))
        for need in ["satiety", "relax", "money"]:
            pressure_peaks[agent_id][need] = max(
                parse_float(row.get(f"{need}_effective_pressure"))
                for row in rows
            )
        for row in rows:
            tool = str(row.get("action_tool") or "")
            if tool:
                behavior_counts[tool] += 1
            if str(row.get("post_content") or ""):
                post_scores.append(parse_float(row.get("post_opinion_index")))
            mediators = json_obj(row.get("mediators"))
            for value in mediators.values():
                mediator_peaks[str(agent_id)] = max(mediator_peaks[str(agent_id)], parse_float(value))
            before = parse_float(row.get("opinion_before"))
            after = parse_float(row.get("opinion_after"))
            if abs(after - before) >= 0.2:
                abnormal_ticks.append({
                    "agent_id": agent_id,
                    "tick": row.get("tick"),
                    "before": before,
                    "after": after,
                    "reason": row.get("opinion_assessment_reason", ""),
                })
            for option, count in json_obj(row.get("opinion_voting_choice_counts")).items():
                vote_totals[str(option)] += int(parse_float(count))
            failed_vote_total += int(parse_float(row.get("opinion_voting_failed_votes")))

    lines = [
        "# 实验运行摘要",
        "",
        "## 配置快照",
        "",
        "```json",
        json.dumps(config_snapshot, ensure_ascii=False, indent=2, default=str),
        "```",
        "",
        "## 保留的连续 opinion（LLM 投票不写回）" if voting_mode else "## 最终观念分布",
        "",
        markdown_table(["agent_id", "final_opinion"], [
            [agent_id, f"{value:.3f}"]
            for agent_id, value in sorted(final_opinions.items())
        ]),
        "",
        "## 有效压力峰值",
        "",
        markdown_table(["agent_id", "satiety", "relax", "money"], [
            [
                agent_id,
                f"{values.get('satiety', 0.0):.3f}",
                f"{values.get('relax', 0.0):.3f}",
                f"{values.get('money', 0.0):.3f}",
            ]
            for agent_id, values in sorted(pressure_peaks.items())
        ]),
        "",
        "## 行为分布",
        "",
        markdown_table(["action_tool", "count"], [
            [tool, count]
            for tool, count in behavior_counts.most_common()
        ]),
        "",
        "## 发帖立场",
        "",
        f"- 发帖样本数：{len(post_scores)}",
        f"- 发帖文本立场均值：{mean(post_scores):.3f}" if post_scores else "- 发帖文本立场均值：无发帖样本",
        "",
        "## LLM 投票汇总",
        "",
        markdown_table(["option", "votes"], [
            [option, count] for option, count in vote_totals.items()
        ]),
        "",
        f"- 失败票数：{failed_vote_total}",
        "",
        "## 输出图表",
        "",
        "- `opinion_trends.svg`",
        "- `opinion_dashboard.html`",
        "- `opinion_distribution.svg`",
        "- `opinion_distribution.csv`",
        "- `opinion_analysis.json`",
        "- `opinion_voting_trends.svg`",
        "- `opinion_voting_polarization_trends.svg`",
        "- `effective_pressure_trends.svg`",
        "- `mediator_peak_trends.svg`",
        "- `memory_event_trends.svg`",
        "- `memory_vector_trends.svg`",
        "- `polarization_report.md`",
        "- `polarization_metrics.csv`",
        "- `polarization_agent_shift.csv`",
        "- `opinion_voting_polarization_report.md`",
        "- `opinion_voting_polarization_metrics.csv`",
        "- `opinion_voting_agent_metrics.csv`",
        "- `opinion_voting_skipped_records.csv`",
        "",
        "## 明显观念变化",
        "",
        markdown_table(["agent_id", "tick", "before", "after", "reason"], [
            [
                item["agent_id"],
                item["tick"],
                f"{item['before']:.3f}",
                f"{item['after']:.3f}",
                str(item["reason"])[:80],
            ]
            for item in abnormal_ticks[:20]
        ]),
        "",
    ]
    return "\n".join(lines)


def write_standard_charts(
    output_dir: Path,
    rows_by_agent: dict[str, list[dict]],
    *,
    voting_rows_by_agent: dict[str, list[dict]] | None = None,
    include_voting_chart: bool = True,
) -> list[Path]:
    """生成标准趋势图，并允许调用方隔离附加投票趋势。"""

    generated = []
    generated.append(write_opinion_trend_svg(output_dir / "opinion_trends.svg", rows_by_agent))
    if include_voting_chart:
        generated.append(
            write_opinion_voting_trend_svg(
                output_dir / "opinion_voting_trends.svg",
                voting_rows_by_agent or rows_by_agent,
            )
        )
    generated.append(write_multi_series_svg(
        output_dir / "effective_pressure_trends.svg",
        rows_by_agent,
        value_getter=lambda row: max(
            parse_float(row.get("satiety_effective_pressure")),
            parse_float(row.get("relax_effective_pressure")),
            parse_float(row.get("money_effective_pressure")),
        ),
        title="Effective Pressure Trends",
        y_min=0.0,
        y_max=1.0,
    ))
    generated.append(write_multi_series_svg(
        output_dir / "mediator_peak_trends.svg",
        rows_by_agent,
        value_getter=lambda row: max((parse_float(v) for v in json_obj(row.get("mediators")).values()), default=0.0),
        title="Mediator Peak Trends",
        y_min=0.0,
        y_max=1.0,
    ))
    generated.append(write_memory_count_trend_svg(
        output_dir / "memory_event_trends.svg",
        rows_by_agent,
        field="active_memory_events",
        title="Active Memory Events",
    ))
    generated.append(write_memory_count_trend_svg(
        output_dir / "memory_vector_trends.svg",
        rows_by_agent,
        field="active_memory_vectors",
        title="Active Memory Vectors",
    ))
    return generated


def write_memory_count_trend_svg(
    output_path: Path,
    rows_by_agent: dict[str, list[dict]],
    *,
    field: str,
    title: str,
) -> Path:
    """绘制每个智能体的活跃记忆数量趋势。"""

    maximum = max(
        (parse_float(row.get(field)) for rows in rows_by_agent.values() for row in rows),
        default=1.0,
    )
    return write_multi_series_svg(
        output_path,
        rows_by_agent,
        value_getter=lambda row: parse_float(row.get(field)),
        title=title,
        y_min=0.0,
        y_max=max(1.0, maximum),
    )


def write_opinion_voting_trend_svg(output_path: Path, rows_by_agent: dict[str, list[dict]]) -> Path:
    """按评测时间步汇总所有智能体的各选项票数。"""

    counts_by_tick: dict[int, Counter] = defaultdict(Counter)
    options: list[str] = []
    seen_options: set[str] = set()
    for rows in rows_by_agent.values():
        for row in rows:
            raw_tick = row.get("opinion_voting_tick")
            if raw_tick in (None, ""):
                continue
            tick = _strict_nonnegative_int(raw_tick)
            requested = _strict_nonnegative_int(row.get("opinion_voting_requested_voters"))
            successful = _strict_nonnegative_int(row.get("opinion_voting_successful_votes"))
            failed = _strict_nonnegative_int(row.get("opinion_voting_failed_votes"))
            counts = json_obj(row.get("opinion_voting_choice_counts"))
            roles = json_obj(row.get("opinion_voting_option_roles"))
            normalized_counts = {
                option: _strict_nonnegative_int(value)
                for option, value in counts.items()
            }
            speech_history_complete = True
            if "opinion_voting_speech_history" in row:
                raw_history = row.get("opinion_voting_speech_history")
                if isinstance(raw_history, list):
                    speech_history = raw_history
                else:
                    try:
                        speech_history = json.loads(str(raw_history))
                    except (TypeError, json.JSONDecodeError):
                        speech_history = None
                speech_history_complete = isinstance(speech_history, list) and bool(speech_history)
            # 趋势图与统计使用相同完整性门槛，避免部分票数造成视觉误导。
            if (
                tick is None
                or requested is None
                or successful is None
                or failed is None
                or requested <= 0
                or successful != requested
                or failed != 0
                or not counts
                or set(counts) != set(roles)
                or not speech_history_complete
                or any(value is None for value in normalized_counts.values())
                or sum(int(value) for value in normalized_counts.values() if value is not None) != successful
            ):
                continue
            for option, count in counts.items():
                option_text = str(option)
                if option_text not in seen_options:
                    seen_options.add(option_text)
                    options.append(option_text)
                counts_by_tick[tick][option_text] += int(normalized_counts[option] or 0)
    rows_by_option = {
        option: [
            {"tick": tick, "vote_count": counts_by_tick[tick].get(option, 0)}
            for tick in sorted(counts_by_tick)
        ]
        for option in options
    }
    max_votes = max(
        (row["vote_count"] for rows in rows_by_option.values() for row in rows),
        default=1,
    )
    return write_multi_series_svg(
        output_path,
        rows_by_option,
        value_getter=lambda row: parse_float(row.get("vote_count")),
        title="LLM Opinion Voting Trends",
        y_min=0.0,
        y_max=float(max(1, max_votes)),
    )


def write_opinion_trend_svg(output_path: Path, rows_by_agent: dict[str, list[dict]]) -> Path:
    return write_multi_series_svg(
        output_path,
        rows_by_agent,
        value_getter=lambda row: parse_float(row.get("opinion")),
        title="Agent Opinion Trends",
        y_min=-1.0,
        y_max=1.0,
    )


def opinion_distribution_categories(options: AnalysisOptions) -> list[dict[str, str]]:
    """返回与现有极化阈值一致、互斥且覆盖全部 opinion 的五档定义。"""

    if not (
        options.neutral_threshold >= 0
        and options.oppose_threshold < -options.neutral_threshold
        and options.support_threshold > options.neutral_threshold
    ):
        raise ValueError(
            "观念分布要求 oppose_threshold < -neutral_threshold <= neutral_threshold < support_threshold。"
        )
    return [
        {"key": "strong_oppose", "label": "明显反对", "color": "#b91c1c"},
        {"key": "lean_oppose", "label": "偏反对", "color": "#f97316"},
        {"key": "neutral", "label": "中立", "color": "#a3a3a3"},
        {"key": "lean_support", "label": "偏支持", "color": "#22c55e"},
        {"key": "strong_support", "label": "明显支持", "color": "#15803d"},
    ]


def opinion_distribution_key(value: float, options: AnalysisOptions) -> str:
    """按照互斥阈值把单个连续观念值归入一档。"""

    if value <= options.oppose_threshold:
        return "strong_oppose"
    if value < -options.neutral_threshold:
        return "lean_oppose"
    if value <= options.neutral_threshold:
        return "neutral"
    if value < options.support_threshold:
        return "lean_support"
    return "strong_support"


def build_opinion_distribution(
    opinions_by_agent: dict[str, float],
    categories: list[dict[str, str]],
    options: AnalysisOptions,
) -> list[dict[str, object]]:
    """计算一个时刻中每档观念的数量和占比。"""

    counts = Counter(opinion_distribution_key(value, options) for value in opinions_by_agent.values())
    total = len(opinions_by_agent)
    return [
        {
            **category,
            "count": counts[category["key"]],
            "share": counts[category["key"]] / total if total else 0.0,
        }
        for category in categories
    ]


def write_opinion_analysis_outputs(
    *,
    run_dir: Path,
    output_dir: Path,
    options: AnalysisOptions,
) -> list[Path]:
    """生成完整观念序列、首末分布和可筛选的交互分析页面。"""

    series_by_agent = read_agent_opinions(run_dir)
    ticks = select_ticks(series_by_agent, allow_incomplete=False)
    if options.data_length is not None:
        ticks = ticks[:options.data_length]
    if not ticks:
        raise ValueError(f"没有可用于观念分布的完整 tick：{run_dir}")
    categories = opinion_distribution_categories(options)
    initial_tick = ticks[0]
    final_tick = ticks[-1]
    stages = []
    for key, label, tick in [
        ("initial", "初始时刻", initial_tick),
        ("final", "结束时刻", final_tick),
    ]:
        opinions = opinions_for_tick(series_by_agent, tick, allow_incomplete=False)
        stages.append({
            "key": key,
            "label": label,
            "tick": tick,
            "n_agents": len(opinions),
            "distribution": build_opinion_distribution(opinions, categories, options),
        })
    payload = {
        "version": 1,
        "thresholds": {
            "oppose_threshold": options.oppose_threshold,
            "neutral_threshold": options.neutral_threshold,
            "support_threshold": options.support_threshold,
        },
        "tick_range": {"start": initial_tick, "end": final_tick, "count": len(ticks)},
        "agent_ids": sorted(series_by_agent),
        "series": [
            {
                "agent_id": agent_id,
                "points": [
                    {"tick": tick, "opinion": series_by_agent[agent_id][tick]}
                    for tick in ticks
                ],
            }
            for agent_id in sorted(series_by_agent)
        ],
        "stages": stages,
    }
    json_path = output_dir / "opinion_analysis.json"
    csv_path = output_dir / "opinion_distribution.csv"
    svg_path = output_dir / "opinion_distribution.svg"
    html_path = output_dir / "opinion_dashboard.html"
    json_path.write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")
    write_opinion_distribution_csv(csv_path, stages)
    write_opinion_distribution_svg(svg_path, stages)
    write_opinion_dashboard_html(html_path, payload)
    return [json_path, csv_path, svg_path, html_path]


def write_opinion_distribution_csv(path: Path, stages: list[dict]) -> None:
    """写出首末时刻各观念档位的数量和占比。"""

    fields = ["stage", "stage_label", "tick", "n_agents", "category", "category_label", "count", "share"]
    with path.open("w", newline="", encoding="utf-8") as f:
        writer = csv.DictWriter(f, fieldnames=fields)
        writer.writeheader()
        for stage in stages:
            for item in stage["distribution"]:
                writer.writerow({
                    "stage": stage["key"],
                    "stage_label": stage["label"],
                    "tick": stage["tick"],
                    "n_agents": stage["n_agents"],
                    "category": item["key"],
                    "category_label": item["label"],
                    "count": item["count"],
                    "share": format_number(item["share"]),
                })


def write_opinion_distribution_svg(path: Path, stages: list[dict]) -> Path:
    """绘制初始与结束时刻的五档观念占比分组柱状图。"""

    width, height = 960, 520
    left, right, top, bottom = 74, 30, 64, 120
    plot_width = width - left - right
    plot_height = height - top - bottom
    categories = stages[0]["distribution"]
    group_width = plot_width / max(1, len(categories))
    bar_width = min(42.0, group_width * 0.28)

    def y_pos(share: float) -> float:
        return top + plot_height * (1.0 - clamp_number(share, 0.0, 1.0))

    lines = [
        f'<svg xmlns="http://www.w3.org/2000/svg" width="{width}" height="{height}" viewBox="0 0 {width} {height}">',
        '<style>.axis{stroke:#333}.grid{stroke:#ddd}.title{font:18px sans-serif;fill:#222}.label{font:13px sans-serif;fill:#333}.tick{font:12px sans-serif;fill:#555}.value{font:11px sans-serif;fill:#222}</style>',
        f'<rect width="{width}" height="{height}" fill="white"/>',
        f'<text class="title" x="{width / 2}" y="32" text-anchor="middle">Initial and Final Opinion Distribution</text>',
    ]
    for share in [0.0, 0.25, 0.5, 0.75, 1.0]:
        y = y_pos(share)
        lines.append(f'<line class="grid" x1="{left}" y1="{y:.2f}" x2="{left + plot_width}" y2="{y:.2f}"/>')
        lines.append(f'<text class="tick" x="{left - 10}" y="{y + 4:.2f}" text-anchor="end">{share:.0%}</text>')
    lines.append(f'<line class="axis" x1="{left}" y1="{top}" x2="{left}" y2="{top + plot_height}"/>')
    lines.append(f'<line class="axis" x1="{left}" y1="{top + plot_height}" x2="{left + plot_width}" y2="{top + plot_height}"/>')
    stage_colors = ["#2563eb", "#dc2626"]
    for category_index, category in enumerate(categories):
        center = left + group_width * (category_index + 0.5)
        lines.append(
            f'<text class="label" x="{center:.2f}" y="{top + plot_height + 28}" text-anchor="middle">'
            f'{escape_text(category["label"])}</text>'
        )
        for stage_index, stage in enumerate(stages):
            item = stage["distribution"][category_index]
            share = float(item["share"])
            x = center + (stage_index - 0.5) * (bar_width + 8) - bar_width / 2
            y = y_pos(share)
            bar_height = top + plot_height - y
            lines.append(
                f'<rect x="{x:.2f}" y="{y:.2f}" width="{bar_width:.2f}" height="{bar_height:.2f}" '
                f'fill="{stage_colors[stage_index]}" rx="2"/>'
            )
            lines.append(
                f'<text class="value" x="{x + bar_width / 2:.2f}" y="{max(top + 12, y - 6):.2f}" text-anchor="middle">'
                f'{share:.1%} ({item["count"]})</text>'
            )
    legend_y = height - 44
    for index, stage in enumerate(stages):
        x = width / 2 - 160 + index * 260
        lines.append(f'<rect x="{x:.2f}" y="{legend_y - 12}" width="16" height="16" fill="{stage_colors[index]}" rx="2"/>')
        lines.append(
            f'<text class="label" x="{x + 24:.2f}" y="{legend_y + 1}">'
            f'{escape_text(stage["label"])} (tick={stage["tick"]}, n={stage["n_agents"]})</text>'
        )
    lines.append("</svg>")
    path.write_text("\n".join(lines), encoding="utf-8")
    return path


def write_opinion_dashboard_html(path: Path, payload: dict) -> Path:
    """生成无需外部依赖的交互式观念分析页面。"""

    payload_json = json.dumps(payload, ensure_ascii=False).replace("</", "<\\/")
    template = """<!doctype html>
<html lang="zh-CN">
<head>
  <meta charset="utf-8" />
  <meta name="viewport" content="width=device-width, initial-scale=1" />
  <title>观念变化分析</title>
  <style>
    :root { font-family: system-ui, -apple-system, "Segoe UI", sans-serif; color: #1f2937; background: #f6f7f4; }
    body { margin: 0; padding: 24px; }
    main { max-width: 1180px; margin: 0 auto; }
    h1 { margin: 0 0 6px; font-size: 24px; }
    h2 { margin: 28px 0 10px; font-size: 17px; }
    .meta { color: #6b7280; font-size: 13px; }
    .toolbar { display: flex; align-items: center; gap: 10px; margin: 18px 0 8px; }
    label { font-size: 13px; font-weight: 600; }
    select { min-width: 260px; border: 1px solid #9ca3af; border-radius: 4px; padding: 7px 10px; background: white; color: #111827; }
    .chart { width: 100%; min-height: 420px; background: white; border: 1px solid #d1d5db; }
    .distribution { display: grid; grid-template-columns: repeat(2, minmax(0, 1fr)); gap: 20px; }
    .stage { background: white; border-top: 3px solid #374151; padding: 14px; }
    .stage h3 { margin: 0 0 12px; font-size: 15px; }
    .row { display: grid; grid-template-columns: 88px 1fr 92px; gap: 10px; align-items: center; margin: 9px 0; font-size: 12px; }
    .track { height: 14px; background: #e5e7eb; }
    .fill { height: 100%; min-width: 0; }
    .number { text-align: right; font-variant-numeric: tabular-nums; }
    @media (max-width: 760px) { body { padding: 14px; } .distribution { grid-template-columns: 1fr; } select { min-width: 0; width: 100%; } }
  </style>
</head>
<body>
<main>
  <h1>观念变化分析</h1>
  <div class="meta" id="meta"></div>
  <div class="toolbar">
    <label for="agent-select">曲线范围</label>
    <select id="agent-select"></select>
  </div>
  <svg id="opinion-chart" class="chart" viewBox="0 0 1000 460" role="img" aria-label="智能体观念变化曲线"></svg>
  <h2>初始与结束时刻观念分布</h2>
  <div id="distribution" class="distribution"></div>
</main>
<script>
const data = __PAYLOAD__;
const palette = ['#2563eb','#dc2626','#16a34a','#9333ea','#ea580c','#0891b2','#be123c','#4d7c0f'];
const select = document.getElementById('agent-select');
const chart = document.getElementById('opinion-chart');
document.getElementById('meta').textContent = `tick ${data.tick_range.start} - ${data.tick_range.end} · ${data.agent_ids.length} 个智能体 · ${data.tick_range.count} 个完整时间点`;

// 下拉框同时支持全体曲线和指定单一智能体。
select.innerHTML = '<option value="__all__">全部智能体</option>' + data.agent_ids.map(id => `<option value="${escapeHtml(id)}">${escapeHtml(id)}</option>`).join('');
select.addEventListener('change', drawChart);

function escapeHtml(value) {
  return String(value).replace(/[&<>"']/g, char => ({'&':'&amp;','<':'&lt;','>':'&gt;','"':'&quot;',"'":'&#39;'}[char]));
}
function linePath(points, x, y) {
  return points.map((point, index) => `${index ? 'L' : 'M'} ${x(point.tick).toFixed(2)} ${y(point.opinion).toFixed(2)}`).join(' ');
}
function drawChart() {
  const selected = select.value;
  const series = selected === '__all__' ? data.series : data.series.filter(item => item.agent_id === selected);
  const left = 70, right = 28, top = 30, bottom = 58, width = 1000, height = 460;
  const plotWidth = width - left - right, plotHeight = height - top - bottom;
  const minTick = data.tick_range.start, maxTick = data.tick_range.end;
  const x = tick => left + (maxTick === minTick ? plotWidth / 2 : (tick - minTick) / (maxTick - minTick) * plotWidth);
  const y = opinion => top + (1 - (opinion + 1) / 2) * plotHeight;
  let html = '<rect width="1000" height="460" fill="white"/>';
  [-1,-0.5,0,0.5,1].forEach(value => {
    const py = y(value);
    html += `<line x1="${left}" y1="${py}" x2="${left + plotWidth}" y2="${py}" stroke="${value === 0 ? '#9ca3af' : '#e5e7eb'}"/>`;
    html += `<text x="${left - 12}" y="${py + 4}" text-anchor="end" font-size="12" fill="#6b7280">${value.toFixed(1)}</text>`;
  });
  html += `<line x1="${left}" y1="${top}" x2="${left}" y2="${top + plotHeight}" stroke="#374151"/>`;
  html += `<line x1="${left}" y1="${top + plotHeight}" x2="${left + plotWidth}" y2="${top + plotHeight}" stroke="#374151"/>`;
  html += `<text x="${left}" y="${height - 20}" font-size="12" fill="#6b7280">tick ${minTick}</text>`;
  html += `<text x="${left + plotWidth}" y="${height - 20}" text-anchor="end" font-size="12" fill="#6b7280">tick ${maxTick}</text>`;
  series.forEach((item, index) => {
    const color = palette[(data.agent_ids.indexOf(item.agent_id) + palette.length) % palette.length];
    html += `<path d="${linePath(item.points, x, y)}" fill="none" stroke="${color}" stroke-width="${selected === '__all__' ? 1.4 : 3}" opacity="${selected === '__all__' ? 0.7 : 1}"><title>${escapeHtml(item.agent_id)}</title></path>`;
    if (selected !== '__all__') {
      item.points.forEach(point => {
        html += `<circle cx="${x(point.tick)}" cy="${y(point.opinion)}" r="3" fill="${color}"><title>${escapeHtml(item.agent_id)} · tick ${point.tick} · opinion ${point.opinion.toFixed(3)}</title></circle>`;
      });
    }
  });
  chart.innerHTML = html;
}

// 分布条直接展示数量和百分比，两个时刻使用同一组分类阈值。
document.getElementById('distribution').innerHTML = data.stages.map(stage => `
  <section class="stage">
    <h3>${escapeHtml(stage.label)} · tick ${stage.tick} · n=${stage.n_agents}</h3>
    ${stage.distribution.map(item => `
      <div class="row">
        <span>${escapeHtml(item.label)}</span>
        <div class="track"><div class="fill" style="width:${(item.share * 100).toFixed(2)}%;background:${item.color}"></div></div>
        <span class="number">${item.count} · ${(item.share * 100).toFixed(1)}%</span>
      </div>`).join('')}
  </section>`).join('');
drawChart();
</script>
</body>
</html>
"""
    path.write_text(template.replace("__PAYLOAD__", payload_json), encoding="utf-8")
    return path


def write_multi_series_svg(
    output_path: Path,
    rows_by_agent: dict[str, list[dict]],
    *,
    value_getter,
    title: str,
    y_min: float,
    y_max: float,
) -> Path:
    width, height = 900, 520
    left, right, top, bottom = 70, 30, 54, 72
    plot_width = width - left - right
    plot_height = height - top - bottom
    colors = ["#2563eb", "#dc2626", "#16a34a", "#9333ea", "#ea580c", "#0891b2", "#be123c", "#4d7c0f"]
    series = {}
    ticks = []
    for agent_id, rows in rows_by_agent.items():
        points = []
        for row in rows:
            tick = int(float(row.get("tick") or 0))
            points.append((tick, value_getter(row)))
            ticks.append(tick)
        series[agent_id] = points
    if not ticks:
        output_path.write_text("<svg xmlns=\"http://www.w3.org/2000/svg\"></svg>", encoding="utf-8")
        return output_path
    min_tick, max_tick = min(ticks), max(ticks)

    def x_pos(tick: int) -> float:
        return scale_value(tick, min_tick, max_tick, left, left + plot_width)

    def y_pos(value: float) -> float:
        bounded = clamp_number(value, y_min, y_max)
        return scale_value(bounded, y_min, y_max, top + plot_height, top)

    lines = [
        f'<svg xmlns="http://www.w3.org/2000/svg" width="{width}" height="{height}" viewBox="0 0 {width} {height}">',
        "<style>.axis{stroke:#333;stroke-width:1}.grid{stroke:#ddd;stroke-width:1}.label{font:14px sans-serif;fill:#222}.tick{font:12px sans-serif;fill:#555}.legend{font:12px sans-serif;fill:#222}</style>",
        f'<rect width="{width}" height="{height}" fill="white"/>',
        f'<text class="label" x="{width / 2}" y="28" text-anchor="middle">{escape_text(title)}</text>',
    ]
    for value in [y_min, (y_min + y_max) / 2, y_max]:
        y = y_pos(value)
        lines.append(f'<line class="grid" x1="{left}" y1="{y:.2f}" x2="{left + plot_width}" y2="{y:.2f}"/>')
        lines.append(f'<text class="tick" x="{left - 10}" y="{y + 4:.2f}" text-anchor="end">{value:.2f}</text>')
    lines.append(f'<line class="axis" x1="{left}" y1="{top}" x2="{left}" y2="{top + plot_height}"/>')
    lines.append(f'<line class="axis" x1="{left}" y1="{top + plot_height}" x2="{left + plot_width}" y2="{top + plot_height}"/>')
    for index, (agent_id, points) in enumerate(sorted(series.items())):
        color = colors[index % len(colors)]
        if not points:
            continue
        point_text = " ".join(f"{x_pos(tick):.2f},{y_pos(value):.2f}" for tick, value in points)
        lines.append(f'<polyline points="{point_text}" fill="none" stroke="{color}" stroke-width="2"/>')
        lines.append(f'<text class="legend" x="{left + 10 + index * 130}" y="{height - 24}" fill="{color}">{escape_text(agent_id)}</text>')
    lines.append("</svg>")
    output_path.write_text("\n".join(lines), encoding="utf-8")
    return output_path


def write_voting_polarization_outputs(
    *,
    rows_by_agent: dict[str, list[dict]],
    output_dir: Path,
    options: AnalysisOptions,
    expected_agent_ids: set[str] | None = None,
) -> list[Path]:
    """生成 LLM 投票模式专用的逐轮指标、个体指标、趋势图和报告。"""

    agent_metrics, skipped_records, observed_ticks = read_voting_agent_metrics(rows_by_agent)
    expected = set(expected_agent_ids or rows_by_agent)
    for tick in observed_ticks:
        recorded_agent_ids = {
            metric.agent_id for metric in agent_metrics if metric.tick == tick
        } | {
            record.agent_id for record in skipped_records if record.tick == tick
        }
        for agent_id in sorted(expected - recorded_agent_ids):
            skipped_records.append(VotingSkippedRecord(
                tick=tick,
                agent_id=agent_id,
                reason="missing_record",
                requested_votes=None,
                successful_votes=None,
                failed_votes=None,
            ))
    metrics_path = output_dir / "opinion_voting_polarization_metrics.csv"
    agent_path = output_dir / "opinion_voting_agent_metrics.csv"
    skipped_path = output_dir / "opinion_voting_skipped_records.csv"
    report_path = output_dir / "opinion_voting_polarization_report.md"
    chart_path = output_dir / "opinion_voting_polarization_trends.svg"
    output_paths = [metrics_path, agent_path, skipped_path, report_path, chart_path]
    write_voting_skipped_records_csv(skipped_path, skipped_records)
    if _all_voting_speech_windows_empty(rows_by_agent):
        # 空发言窗口属于正常无数据状态，保留审计记录但不计算极化指标。
        write_voting_metrics_csv(metrics_path, [])
        write_voting_agent_metrics_csv(agent_path, agent_metrics)
        write_voting_polarization_chart(chart_path, [])
        write_voting_no_data_report(
            report_path,
            agent_metrics,
            skipped_records,
            metrics_path,
            agent_path,
            skipped_path,
            reason="全部 posthoc voting 记录的线上发言窗口均为空。",
        )
        return output_paths

    # 只使用请求全部成功且字段关系完整的记录，不用部分票数替代完整投票。
    ticks = sorted({metric.tick for metric in agent_metrics})
    if options.data_length is not None:
        ticks = ticks[:options.data_length]
        tick_set = set(ticks)
        agent_metrics = [metric for metric in agent_metrics if metric.tick in tick_set]
        skipped_records = [record for record in skipped_records if record.tick in tick_set]
        write_voting_skipped_records_csv(skipped_path, skipped_records)
    if not ticks:
        write_voting_metrics_csv(metrics_path, [])
        write_voting_agent_metrics_csv(agent_path, [])
        write_voting_polarization_chart(chart_path, [])
        write_voting_no_data_report(
            report_path,
            [],
            skipped_records,
            metrics_path,
            agent_path,
            skipped_path,
            reason="没有完整投票记录可进入统计。",
        )
        return output_paths
    tick_metrics = compute_voting_tick_metrics(agent_metrics, ticks)
    write_voting_metrics_csv(metrics_path, tick_metrics)
    write_voting_agent_metrics_csv(agent_path, agent_metrics)
    write_voting_polarization_chart(chart_path, tick_metrics)
    try:
        baseline_ticks, final_ticks = split_windows(
            ticks,
            options.voting_baseline_window_size,
            options.voting_final_window_size,
            series_name="llm_voting 完整评测轮次",
        )
    except ValueError as exc:
        write_voting_no_data_report(
            report_path,
            agent_metrics,
            skipped_records,
            metrics_path,
            agent_path,
            skipped_path,
            reason=str(exc),
        )
        return output_paths
    stats = summarize_voting_windows(tick_metrics, baseline_ticks, final_ticks)
    ci_low, ci_high = bootstrap_voting_polarization_delta(
        agent_metrics,
        baseline_ticks,
        final_ticks,
        samples=options.bootstrap_samples,
        seed=options.seed,
    )

    write_voting_polarization_report(
        report_path=report_path,
        metrics_path=metrics_path,
        agent_path=agent_path,
        options=options,
        ticks=ticks,
        baseline_ticks=baseline_ticks,
        final_ticks=final_ticks,
        stats=stats,
        ci_low=ci_low,
        ci_high=ci_high,
        agent_metrics=agent_metrics,
        tick_metrics=tick_metrics,
        skipped_records=skipped_records,
        skipped_path=skipped_path,
    )
    return output_paths


def _all_voting_speech_windows_empty(rows_by_agent: dict[str, list[dict]]) -> bool:
    """仅识别明确记录了空 speech history 的 posthoc 投票数据。"""

    voting_rows = [
        row
        for rows in rows_by_agent.values()
        for row in rows
        if row.get("opinion_voting_tick") not in (None, "")
    ]
    if not voting_rows:
        return False
    for row in voting_rows:
        if "opinion_voting_speech_history" not in row:
            return False
        raw_history = row.get("opinion_voting_speech_history")
        if isinstance(raw_history, list):
            history = raw_history
        else:
            try:
                history = json.loads(str(raw_history))
            except (TypeError, json.JSONDecodeError):
                return False
        if not isinstance(history, list) or history:
            return False
    return True


def write_voting_no_data_report(
    report_path: Path,
    agent_metrics: list[VotingAgentMetric],
    skipped_records: list[VotingSkippedRecord],
    metrics_path: Path,
    agent_path: Path,
    skipped_path: Path,
    *,
    reason: str,
) -> None:
    """为没有足够完整记录的运行生成可审计报告。"""

    ticks = sorted({metric.tick for metric in agent_metrics})
    agent_ids = sorted(
        {metric.agent_id for metric in agent_metrics}
        | {record.agent_id for record in skipped_records}
    )
    lines = [
        "# LLM 投票极化统计报告",
        "",
        "## 数据状态",
        "",
        "- 状态：完整数据不足，未进行首末窗口比较",
        f"- 原因：{reason}",
        f"- 纳入统计的完整记录数：{len(agent_metrics)}",
        f"- 跳过记录数：{len(skipped_records)}",
        f"- 智能体数量：{len(agent_ids)}",
        f"- 完整记录所在 tick：{', '.join(str(tick) for tick in ticks) if ticks else '无'}",
        "",
        "未计算总体投票指标、bootstrap 置信区间或极化判定。",
        "",
        "## 输出文件",
        "",
        f"- `{metrics_path}`",
        f"- `{agent_path}`（仅包含完整投票记录）",
        f"- `{skipped_path}`（记录全部跳过原因）",
        "",
    ]
    report_path.write_text("\n".join(lines), encoding="utf-8")


def _strict_nonnegative_int(value: object) -> int | None:
    """只接受有限、非负且没有小数部分的整数值。"""

    try:
        number = float(str(value))
    except (TypeError, ValueError):
        return None
    if not math.isfinite(number) or number < 0 or not number.is_integer():
        return None
    return int(number)


def read_voting_agent_metrics(
    rows_by_agent: dict[str, list[dict]],
) -> tuple[list[VotingAgentMetric], list[VotingSkippedRecord], set[int]]:
    """只返回完整投票指标，并为所有未纳入记录保存精确原因。"""

    metrics: list[VotingAgentMetric] = []
    skipped: list[VotingSkippedRecord] = []
    observed_ticks: set[int] = set()
    valid_roles = {"support", "oppose", "unknown"}
    keyed_rows: list[tuple[str, int | None, dict]] = []
    key_counts: Counter = Counter()
    for agent_id, rows in sorted(rows_by_agent.items()):
        for row in rows:
            raw_tick = row.get("opinion_voting_tick")
            if raw_tick in (None, ""):
                continue
            tick = _strict_nonnegative_int(raw_tick)
            keyed_rows.append((agent_id, tick, row))
            if tick is not None:
                observed_ticks.add(tick)
                key_counts[(agent_id, tick)] += 1

    previous_valid_by_agent: dict[str, VotingAgentMetric] = {}
    for agent_id, tick, row in sorted(
        keyed_rows,
        key=lambda item: (item[0], item[1] if item[1] is not None else -1),
    ):
        requested_votes = _strict_nonnegative_int(row.get("opinion_voting_requested_voters"))
        successful_votes = _strict_nonnegative_int(row.get("opinion_voting_successful_votes"))
        failed_votes = _strict_nonnegative_int(row.get("opinion_voting_failed_votes"))

        def skip(reason: str) -> None:
            skipped.append(VotingSkippedRecord(
                tick=tick,
                agent_id=agent_id,
                reason=reason,
                requested_votes=requested_votes,
                successful_votes=successful_votes,
                failed_votes=failed_votes,
            ))

        if tick is None:
            skip("invalid_tick")
            continue
        if key_counts[(agent_id, tick)] > 1:
            skip("duplicate_record")
            continue
        if requested_votes is None or successful_votes is None or failed_votes is None:
            skip("invalid_vote_totals")
            continue
        skipped_reason = str(row.get("opinion_voting_skipped_reason") or "").strip()
        if "opinion_voting_speech_history" in row:
            raw_history = row.get("opinion_voting_speech_history")
            if isinstance(raw_history, list):
                speech_history = raw_history
            else:
                try:
                    speech_history = json.loads(str(raw_history))
                except (TypeError, json.JSONDecodeError):
                    skip("invalid_speech_history")
                    continue
            if not isinstance(speech_history, list):
                skip("invalid_speech_history")
                continue
            if not speech_history:
                skip(skipped_reason or "no_online_speech_in_window")
                continue
        if requested_votes == 0:
            skip(skipped_reason or "no_requested_votes")
            continue
        if requested_votes != successful_votes + failed_votes:
            skip("inconsistent_vote_totals")
            continue
        if failed_votes > 0 or successful_votes != requested_votes:
            skip("incomplete_votes")
            continue

        counts = json_obj(row.get("opinion_voting_choice_counts"))
        roles = json_obj(row.get("opinion_voting_option_roles"))
        if not counts or not roles:
            skip("missing_counts_or_roles")
            continue
        if set(counts) != set(roles):
            skip("count_role_options_mismatch")
            continue
        if set(roles.values()) - valid_roles:
            skip("invalid_option_role")
            continue
        if "support" not in roles.values() or "oppose" not in roles.values():
            skip("missing_direction_role")
            continue
        normalized_counts = {
            option: _strict_nonnegative_int(value)
            for option, value in counts.items()
        }
        if any(value is None for value in normalized_counts.values()):
            skip("invalid_choice_count")
            continue
        complete_counts = {
            option: int(value)
            for option, value in normalized_counts.items()
            if value is not None
        }
        if sum(complete_counts.values()) != successful_votes:
            skip("choice_count_total_mismatch")
            continue

        role_counts = {
            role: sum(count for option, count in complete_counts.items() if roles[option] == role)
            for role in valid_roles
        }
        divisor = float(successful_votes)
        support_share = role_counts["support"] / divisor
        oppose_share = role_counts["oppose"] / divisor
        unknown_share = role_counts["unknown"] / divisor
        known_share = support_share + oppose_share
        decisiveness = abs(support_share - oppose_share) / known_share if known_share else 0.0
        stance_result = classify_voting_stance(
            requested_voters=requested_votes,
            successful_votes=successful_votes,
            choice_counts=complete_counts,
            option_roles=roles,
        )
        stance_valid = int(bool(stance_result["stance_valid"]))
        stance = str(stance_result["stance"])
        previous_valid = previous_valid_by_agent.get(agent_id)
        metric = VotingAgentMetric(
            tick=tick,
            agent_id=agent_id,
            requested_votes=requested_votes,
            successful_votes=successful_votes,
            failed_votes=failed_votes,
            skipped_reason=skipped_reason,
            support_share=support_share,
            oppose_share=oppose_share,
            unknown_share=unknown_share,
            known_share=known_share,
            decisiveness=decisiveness,
            stance=stance,
            stance_valid=stance_valid,
            stance_agreement=float(stance_result["stance_agreement"]),
            stance_direction_margin=float(stance_result["stance_direction_margin"]),
            stance_success_rate=float(stance_result["stance_success_rate"]),
            previous_valid_tick=previous_valid.tick if previous_valid is not None else 0,
            previous_valid_stance=previous_valid.stance if previous_valid is not None else "",
            stance_changed=int(
                bool(stance_valid and previous_valid is not None and previous_valid.stance != stance)
            ),
        )
        metrics.append(metric)
        if stance_valid:
            previous_valid_by_agent[agent_id] = metric
    return metrics, skipped, observed_ticks


def compute_voting_tick_metrics(
    agent_metrics: list[VotingAgentMetric],
    ticks: list[int],
) -> list[VotingTickMetric]:
    """对每轮投票先按智能体归一化，再计算总体极化。"""

    by_tick: dict[int, list[VotingAgentMetric]] = defaultdict(list)
    for metric in agent_metrics:
        by_tick[metric.tick].append(metric)
    out = []
    for tick in ticks:
        rows = by_tick[tick]
        valid = [row for row in rows if row.successful_votes > 0]
        if not valid:
            raise ValueError(f"投票 tick {tick} 没有成功票。")
        support_share = mean(row.support_share for row in valid)
        oppose_share = mean(row.oppose_share for row in valid)
        unknown_share = mean(row.unknown_share for row in valid)
        known_share = support_share + oppose_share
        decisiveness = mean(row.decisiveness for row in valid)
        camp_balance = 4.0 * (support_share / known_share) * (oppose_share / known_share) if known_share else 0.0
        polarization_index = camp_balance * known_share * decisiveness
        requested_votes = sum(row.requested_votes for row in rows)
        successful_votes = sum(row.successful_votes for row in rows)
        failed_votes = sum(row.failed_votes for row in rows)
        stance_counts = Counter(row.stance for row in rows)
        stance_divisor = float(len(rows)) if rows else 1.0
        comparable = [row for row in rows if row.stance_valid and row.previous_valid_tick > 0]
        changed = [row for row in comparable if row.stance_changed]
        out.append(VotingTickMetric(
            tick=tick,
            n_agents=len(rows),
            valid_agents=len(valid),
            requested_votes=requested_votes,
            successful_votes=successful_votes,
            failed_votes=failed_votes,
            success_rate=successful_votes / requested_votes if requested_votes else 0.0,
            support_share=support_share,
            oppose_share=oppose_share,
            unknown_share=unknown_share,
            known_share=known_share,
            decisiveness=decisiveness,
            camp_balance=camp_balance,
            polarization_index=polarization_index,
            stance_valid_agents=len(rows) - stance_counts[VOTING_STANCE_INVALID],
            stance_invalid_agents=stance_counts[VOTING_STANCE_INVALID],
            stance_support_count=stance_counts[VOTING_ROLE_SUPPORT],
            stance_oppose_count=stance_counts[VOTING_ROLE_OPPOSE],
            stance_unknown_count=stance_counts[VOTING_ROLE_UNKNOWN],
            stance_support_share=stance_counts[VOTING_ROLE_SUPPORT] / stance_divisor,
            stance_oppose_share=stance_counts[VOTING_ROLE_OPPOSE] / stance_divisor,
            stance_unknown_share=stance_counts[VOTING_ROLE_UNKNOWN] / stance_divisor,
            stance_invalid_share=stance_counts[VOTING_STANCE_INVALID] / stance_divisor,
            stance_comparable_agents=len(comparable),
            stance_changed_agents=len(changed),
            stance_unchanged_agents=len(comparable) - len(changed),
            stance_changed_share=len(changed) / len(comparable) if comparable else 0.0,
            stance_changed_to_support=sum(row.stance == VOTING_ROLE_SUPPORT for row in changed),
            stance_changed_to_oppose=sum(row.stance == VOTING_ROLE_OPPOSE for row in changed),
            stance_changed_to_unknown=sum(row.stance == VOTING_ROLE_UNKNOWN for row in changed),
        ))
    return out


def summarize_voting_windows(
    metrics: list[VotingTickMetric],
    baseline_ticks: list[int],
    final_ticks: list[int],
) -> dict[str, float]:
    by_tick = {metric.tick: metric for metric in metrics}
    names = [
        "success_rate", "support_share", "oppose_share", "unknown_share",
        "known_share", "decisiveness", "camp_balance", "polarization_index",
    ]
    out: dict[str, float] = {}
    for name in names:
        baseline = metric_mean([by_tick[tick] for tick in baseline_ticks], name)
        final = metric_mean([by_tick[tick] for tick in final_ticks], name)
        out[f"baseline_{name}"] = baseline
        out[f"final_{name}"] = final
        out[f"delta_{name}"] = final - baseline
    return out


def voting_window_index(
    by_agent_tick: dict[tuple[str, int], VotingAgentMetric],
    sampled_agents: list[str],
    ticks: list[int],
) -> float:
    values = []
    for tick in ticks:
        rows = [by_agent_tick[(agent_id, tick)] for agent_id in sampled_agents if (agent_id, tick) in by_agent_tick]
        valid = [row for row in rows if row.successful_votes > 0]
        if not valid:
            continue
        support = mean(row.support_share for row in valid)
        oppose = mean(row.oppose_share for row in valid)
        known = support + oppose
        decisiveness = mean(row.decisiveness for row in valid)
        balance = 4.0 * (support / known) * (oppose / known) if known else 0.0
        values.append(balance * known * decisiveness)
    return mean(values) if values else 0.0


def bootstrap_voting_polarization_delta(
    agent_metrics: list[VotingAgentMetric],
    baseline_ticks: list[int],
    final_ticks: list[int],
    *,
    samples: int,
    seed: int,
) -> tuple[float | None, float | None]:
    if samples == 0:
        return None, None
    by_agent_tick = {(metric.agent_id, metric.tick): metric for metric in agent_metrics}
    agent_ids = sorted({metric.agent_id for metric in agent_metrics})
    eligible = [
        agent_id for agent_id in agent_ids
        if any((agent_id, tick) in by_agent_tick for tick in baseline_ticks)
        and any((agent_id, tick) in by_agent_tick for tick in final_ticks)
    ]
    if not eligible:
        return None, None
    rng = random.Random(seed)
    deltas = []
    for _ in range(samples):
        sampled = [eligible[rng.randrange(len(eligible))] for _ in eligible]
        baseline = voting_window_index(by_agent_tick, sampled, baseline_ticks)
        final = voting_window_index(by_agent_tick, sampled, final_ticks)
        deltas.append(final - baseline)
    deltas.sort()
    low_index = max(0, int(0.025 * (len(deltas) - 1)))
    high_index = min(len(deltas) - 1, int(0.975 * (len(deltas) - 1)))
    return deltas[low_index], deltas[high_index]


def write_voting_metrics_csv(path: Path, metrics: list[VotingTickMetric]) -> None:
    fields = list(VotingTickMetric.__dataclass_fields__)
    with path.open("w", newline="", encoding="utf-8") as f:
        writer = csv.DictWriter(f, fieldnames=fields)
        writer.writeheader()
        for metric in metrics:
            writer.writerow({field: format_number(getattr(metric, field)) for field in fields})


def write_voting_agent_metrics_csv(path: Path, metrics: list[VotingAgentMetric]) -> None:
    fields = list(VotingAgentMetric.__dataclass_fields__)
    with path.open("w", newline="", encoding="utf-8") as f:
        writer = csv.DictWriter(f, fieldnames=fields)
        writer.writeheader()
        for metric in metrics:
            writer.writerow({field: format_number(getattr(metric, field)) for field in fields})


def write_voting_skipped_records_csv(path: Path, records: list[VotingSkippedRecord]) -> None:
    """保存未进入总体统计的投票记录及其跳过原因。"""

    fields = list(VotingSkippedRecord.__dataclass_fields__)
    with path.open("w", newline="", encoding="utf-8") as f:
        writer = csv.DictWriter(f, fieldnames=fields)
        writer.writeheader()
        for record in sorted(records, key=lambda item: (item.tick is None, item.tick or -1, item.agent_id)):
            writer.writerow({field: format_number(getattr(record, field)) for field in fields})


def build_voting_stance_report_sections(
    agent_metrics: list[VotingAgentMetric],
    tick_metrics: list[VotingTickMetric],
) -> list[str]:
    """生成可直接并入投票极化报告的立场变化章节。"""

    if not tick_metrics:
        raise ValueError("没有可生成立场趋势报告的投票轮次。")
    final = tick_metrics[-1]
    changed_transitions = [row for row in agent_metrics if row.stance_changed]
    transition_counts = Counter(
        f"{row.previous_valid_stance}->{row.stance}"
        for row in changed_transitions
    )
    metrics_by_agent: dict[str, list[VotingAgentMetric]] = defaultdict(list)
    for metric in agent_metrics:
        metrics_by_agent[metric.agent_id].append(metric)
    changes_by_agent = Counter(row.agent_id for row in changed_transitions)

    agent_rows = []
    for agent_id, rows in sorted(metrics_by_agent.items()):
        ordered = sorted(rows, key=lambda item: item.tick)
        valid = [row for row in ordered if row.stance_valid]
        first = valid[0] if valid else None
        last = valid[-1] if valid else None
        agent_rows.append([
            agent_id,
            voting_stance_label(first.stance) if first else "无有效评测",
            voting_stance_label(last.stance) if last else "无有效评测",
            len(valid),
            len(ordered) - len(valid),
            changes_by_agent[agent_id],
        ])

    tick_rows = [
        [
            metric.tick,
            metric.stance_support_count,
            metric.stance_oppose_count,
            metric.stance_unknown_count,
            metric.stance_invalid_agents,
            metric.stance_changed_agents,
            format_number(metric.stance_changed_share),
        ]
        for metric in tick_metrics
    ]
    transition_rows = [
        [voting_transition_label(key), count]
        for key, count in sorted(transition_counts.items())
    ] or [["无立场变化", 0]]
    lines = [
        "## 立场分类与变化趋势", "",
        "### 分类规则", "",
        "- 成功票比例低于 0.80 时记为评测无效，不算作立场未知。",
        "- 未知票比例不低于 0.50 时记为未知。",
        "- 支持或反对票比例不低于 0.60，且相对另一方向领先不低于 0.20，才判为对应立场。",
        "- 其他有效结果记为未知。", "",
        "### 每轮全体分布", "",
        markdown_table(
            ["tick", "支持", "反对", "未知", "评测无效", "发生变化", "变化比例"],
            tick_rows,
        ), "",
        "### 最后一轮", "",
        f"- 时间步：{final.tick}",
        f"- 支持：{final.stance_support_count}/{final.n_agents}（{format_number(final.stance_support_share)}）",
        f"- 反对：{final.stance_oppose_count}/{final.n_agents}（{format_number(final.stance_oppose_share)}）",
        f"- 未知：{final.stance_unknown_count}/{final.n_agents}（{format_number(final.stance_unknown_share)}）",
        f"- 评测无效：{final.stance_invalid_agents}/{final.n_agents}（{format_number(final.stance_invalid_share)}）", "",
        "### 立场转移", "",
        f"- 可比较转移次数：{sum(metric.stance_comparable_agents for metric in tick_metrics)}",
        f"- 实际变化次数：{len(changed_transitions)}", "",
        markdown_table(["变化方向", "次数"], transition_rows), "",
        "### 逐智能体摘要", "",
        markdown_table(
            ["agent_id", "首次有效立场", "最后有效立场", "有效轮次", "无效轮次", "变化次数"],
            agent_rows,
        ), "",
    ]
    return lines


def voting_stance_label(stance: str) -> str:
    return {
        VOTING_ROLE_SUPPORT: "支持",
        VOTING_ROLE_OPPOSE: "反对",
        VOTING_ROLE_UNKNOWN: "未知",
        VOTING_STANCE_INVALID: "评测无效",
    }.get(stance, stance)


def voting_transition_label(value: str) -> str:
    before, separator, after = value.partition("->")
    if not separator:
        return value
    return f"{voting_stance_label(before)} -> {voting_stance_label(after)}"


def write_voting_polarization_chart(path: Path, metrics: list[VotingTickMetric]) -> Path:
    series = {
        field: [{"tick": metric.tick, "value": getattr(metric, field)} for metric in metrics]
        for field in [
            "polarization_index",
            "stance_support_share",
            "stance_oppose_share",
            "stance_unknown_share",
            "stance_invalid_share",
        ]
    }
    return write_multi_series_svg(
        path,
        series,
        value_getter=lambda row: parse_float(row.get("value")),
        title="LLM Voting Polarization Trends",
        y_min=0.0,
        y_max=1.0,
    )


def write_voting_polarization_report(
    *,
    report_path: Path,
    metrics_path: Path,
    agent_path: Path,
    options: AnalysisOptions,
    ticks: list[int],
    baseline_ticks: list[int],
    final_ticks: list[int],
    stats: dict[str, float],
    ci_low: float | None,
    ci_high: float | None,
    agent_metrics: list[VotingAgentMetric],
    tick_metrics: list[VotingTickMetric],
    skipped_records: list[VotingSkippedRecord],
    skipped_path: Path,
) -> None:
    final_known = stats["final_known_share"]
    support_known = stats["final_support_share"] / final_known if final_known else 0.0
    oppose_known = stats["final_oppose_share"] / final_known if final_known else 0.0
    side_pass = min(support_known, oppose_known) >= options.min_side_share
    known_pass = final_known >= options.min_voting_known_share
    decisiveness_pass = stats["final_decisiveness"] >= options.min_voting_decisiveness
    delta_pass = stats["delta_polarization_index"] >= options.min_voting_polarization_delta
    unknown_pass = stats["delta_unknown_share"] <= 0.0
    ci_pass = ci_low is not None and ci_low > 0.0
    strict_pass = side_pass and known_pass and decisiveness_pass and delta_pass and unknown_pass and ci_pass
    names = [
        "success_rate", "support_share", "oppose_share", "unknown_share",
        "known_share", "decisiveness", "camp_balance", "polarization_index",
    ]
    lines = [
        "# LLM 投票极化统计报告", "", "## 数据范围", "",
        f"- 投票轮次数：{len(ticks)}",
        f"- 投票 tick：{', '.join(str(tick) for tick in ticks)}",
        f"- 基线轮次：{', '.join(str(tick) for tick in baseline_ticks)}",
        f"- 末端轮次：{', '.join(str(tick) for tick in final_ticks)}",
        f"- 纳入统计的完整记录数：{len(agent_metrics)}",
        f"- 跳过记录数：{len(skipped_records)}", "",
        "## 指标定义", "",
        "- 只有请求票数全部成功、票数之和一致且选项角色完整的记录进入统计。",
        "- 失败、缺失、重复或结构不完整的记录写入跳过审计表。",
        "- 所有份额先在单个智能体内部按成功票归一化，再对智能体取平均。",
        "- `camp_balance = 4 × support_known × oppose_known`。",
        "- `polarization_index = camp_balance × known_share × decisiveness`。", "",
        "## 前后窗口统计", "",
        markdown_table(["指标", "基线窗口均值", "末端窗口均值", "变化"], [metric_row(stats, name) for name in names]),
        "", "## bootstrap 检验", "",
        f"- 投票极化指数变化 bootstrap 95% CI：[{format_number(ci_low)}, {format_number(ci_high)}]", "",
        "## 判定标准", "",
        markdown_table(["标准", "是否满足"], [
            [f"已知票中的支持、反对份额均不低于 {options.min_side_share}", yes_no(side_pass)],
            [f"有效立场覆盖率不低于 {options.min_voting_known_share}", yes_no(known_pass)],
            [f"个体明确度不低于 {options.min_voting_decisiveness}", yes_no(decisiveness_pass)],
            [f"投票极化指数增长不低于 {options.min_voting_polarization_delta}", yes_no(delta_pass)],
            ["未知份额没有上升", yes_no(unknown_pass)],
            ["bootstrap 置信区间下界大于 0", yes_no(ci_pass)],
        ]),
        "", f"**严格判定：{yes_no(strict_pass)}**", "",
    ]
    lines.extend(build_voting_stance_report_sections(agent_metrics, tick_metrics))
    lines.extend([
        "## 输出文件", "",
        f"- `{metrics_path}`", f"- `{agent_path}`", f"- `{skipped_path}`", "",
    ])
    report_path.write_text("\n".join(lines), encoding="utf-8")


def write_polarization_outputs(run_dir: Path, output_dir: Path, options: AnalysisOptions) -> list[Path]:
    series_by_agent = read_agent_opinions(run_dir)
    ticks = select_ticks(series_by_agent, allow_incomplete=options.allow_incomplete_ticks)
    if options.data_length is not None:
        # 统计范围按时间顺序从开头截取，便于观察实验早期变化。
        ticks = ticks[:options.data_length]
    if not ticks:
        raise ValueError(f"没有可分析的完整 tick：{run_dir}")

    metrics = [
        compute_tick_metric(
            tick,
            opinions_for_tick(series_by_agent, tick, allow_incomplete=options.allow_incomplete_ticks),
            support_threshold=options.support_threshold,
            oppose_threshold=options.oppose_threshold,
            neutral_threshold=options.neutral_threshold,
            extreme_threshold=options.extreme_threshold,
        )
        for tick in ticks
    ]
    metrics_path = output_dir / "polarization_metrics.csv"
    shifts_path = output_dir / "polarization_agent_shift.csv"
    report_path = output_dir / "polarization_report.md"
    # 短运行仍可输出逐 tick 指标，但不放宽首尾窗口必须互斥的要求。
    write_metrics_csv(metrics_path, metrics)
    write_agent_shift_csv(shifts_path, [])
    baseline_ticks, final_ticks = split_windows(
        ticks,
        options.opinion_baseline_window_size,
        options.opinion_final_window_size,
        series_name="连续 opinion tick",
    )
    shifts = compute_agent_shifts(series_by_agent, baseline_ticks, final_ticks)
    stats = summarize_windows(metrics, baseline_ticks, final_ticks)
    sign_test_p = one_sided_sign_test([shift.delta_abs_opinion for shift in shifts])
    ci_low, ci_high = bootstrap_mean_ci(
        [shift.delta_abs_opinion for shift in shifts],
        samples=options.bootstrap_samples,
        seed=options.seed,
    )

    write_agent_shift_csv(shifts_path, shifts)
    write_polarization_report(
        report_path=report_path,
        run_dir=run_dir,
        metrics_path=metrics_path,
        shifts_path=shifts_path,
        options=options,
        series_by_agent=series_by_agent,
        ticks=ticks,
        baseline_ticks=baseline_ticks,
        final_ticks=final_ticks,
        stats=stats,
        sign_test_p=sign_test_p,
        ci_low=ci_low,
        ci_high=ci_high,
    )
    return [metrics_path, shifts_path, report_path]


def read_agent_opinions(run_dir: Path) -> dict[str, dict[int, float]]:
    series_by_agent: dict[str, dict[int, float]] = {}
    for csv_path in history_csv_paths(run_dir):
        agent_id = csv_path.stem
        series: dict[int, float] = {}
        with csv_path.open("r", newline="", encoding="utf-8-sig") as f:
            reader = csv.DictReader(f)
            fieldnames = reader.fieldnames or []
            missing = [column for column in REQUIRED_HISTORY_COLUMNS if column not in fieldnames]
            if missing:
                raise ValueError(f"{csv_path} 缺少列 {missing}；实际表头：{fieldnames}")
            for row_number, row in enumerate(reader, start=2):
                tick = parse_tick(row.get("tick"), csv_path, row_number)
                opinion = parse_opinion(row.get("opinion"), csv_path, row_number)
                if tick in series:
                    raise ValueError(f"{csv_path} 存在重复 tick：{tick}")
                series[tick] = opinion
        if series:
            series_by_agent[agent_id] = series
    if not series_by_agent:
        raise ValueError(f"没有读取到智能体 opinion 数据：{run_dir}")
    return series_by_agent


def parse_tick(value: object, csv_path: Path, row_number: int) -> int:
    try:
        return int(float(str(value)))
    except (TypeError, ValueError) as exc:
        raise ValueError(f"{csv_path} 第 {row_number} 行 tick 不是数值：{value!r}") from exc


def parse_opinion(value: object, csv_path: Path, row_number: int) -> float:
    try:
        opinion = float(str(value))
    except (TypeError, ValueError) as exc:
        raise ValueError(f"{csv_path} 第 {row_number} 行 opinion 不是数值：{value!r}") from exc
    if not math.isfinite(opinion):
        raise ValueError(f"{csv_path} 第 {row_number} 行 opinion 不是有限数值。")
    if opinion < -1.0 or opinion > 1.0:
        raise ValueError(f"{csv_path} 第 {row_number} 行 opinion 超出 [-1, 1]：{opinion}")
    return opinion


def select_ticks(series_by_agent: dict[str, dict[int, float]], *, allow_incomplete: bool) -> list[int]:
    tick_sets = [set(series.keys()) for series in series_by_agent.values()]
    if allow_incomplete:
        return sorted(set().union(*tick_sets))
    return sorted(set.intersection(*tick_sets))


def opinions_for_tick(
    series_by_agent: dict[str, dict[int, float]],
    tick: int,
    *,
    allow_incomplete: bool,
) -> dict[str, float]:
    opinions = {
        agent_id: series[tick]
        for agent_id, series in series_by_agent.items()
        if tick in series
    }
    if not allow_incomplete and len(opinions) != len(series_by_agent):
        raise ValueError(f"tick {tick} 缺少智能体数据。")
    return opinions


def compute_tick_metric(
    tick: int,
    opinions_by_agent: dict[str, float],
    *,
    support_threshold: float,
    oppose_threshold: float,
    neutral_threshold: float,
    extreme_threshold: float,
) -> TickMetric:
    values = list(opinions_by_agent.values())
    if not values:
        raise ValueError(f"tick {tick} 没有 opinion 数据。")

    n_agents = len(values)
    mean_opinion = sum(values) / n_agents
    std_population = math.sqrt(sum((value - mean_opinion) ** 2 for value in values) / n_agents)
    abs_values = [abs(value) for value in values]
    support_values = [value for value in values if value >= support_threshold]
    oppose_values = [value for value in values if value <= oppose_threshold]
    support_count = len(support_values)
    oppose_count = len(oppose_values)
    neutral_count = sum(1 for value in values if abs(value) <= neutral_threshold)
    extreme_count = sum(1 for value in values if abs(value) >= extreme_threshold)
    support_share = support_count / n_agents
    oppose_share = oppose_count / n_agents
    neutral_share = neutral_count / n_agents
    extreme_share = extreme_count / n_agents
    two_side_share = min(support_share, oppose_share)
    pairwise_distance = mean_pairwise_abs_distance(values)
    camp_gap = None
    if support_values and oppose_values:
        camp_gap = sum(support_values) / len(support_values) - sum(oppose_values) / len(oppose_values)
    polarization_index = pairwise_distance * two_side_share * (1.0 - neutral_share)

    return TickMetric(
        tick=tick,
        n_agents=n_agents,
        mean_opinion=mean_opinion,
        median_opinion=float(median(values)),
        std_population=std_population,
        mean_abs_opinion=sum(abs_values) / n_agents,
        median_abs_opinion=float(median(abs_values)),
        pairwise_mean_abs_distance=pairwise_distance,
        min_opinion=min(values),
        max_opinion=max(values),
        range_opinion=max(values) - min(values),
        support_count=support_count,
        oppose_count=oppose_count,
        neutral_count=neutral_count,
        extreme_count=extreme_count,
        support_share=support_share,
        oppose_share=oppose_share,
        neutral_share=neutral_share,
        extreme_share=extreme_share,
        two_side_share=two_side_share,
        camp_gap=camp_gap,
        polarization_index=polarization_index,
    )


def mean_pairwise_abs_distance(values: list[float]) -> float:
    if len(values) < 2:
        return 0.0
    total = 0.0
    count = 0
    for left_index, left in enumerate(values):
        for right in values[left_index + 1:]:
            total += abs(left - right)
            count += 1
    return total / count


def split_windows(
    ticks: list[int],
    baseline_window_size: int,
    final_window_size: int,
    *,
    series_name: str,
) -> tuple[list[int], list[int]]:
    """按独立首尾长度切分，并拒绝任何重叠窗口。"""

    if baseline_window_size < 1 or final_window_size < 1:
        raise ValueError(f"{series_name} 的基线和末端窗口长度必须大于等于 1。")
    required = baseline_window_size + final_window_size
    if len(ticks) < required:
        raise ValueError(
            f"{series_name} 共有 {len(ticks)} 个可用点，基线窗口 {baseline_window_size} 个、"
            f"末端窗口 {final_window_size} 个，共需要 {required} 个不重叠点。"
        )
    baseline_ticks = ticks[:baseline_window_size]
    final_ticks = ticks[-final_window_size:]
    if set(baseline_ticks) & set(final_ticks):
        raise ValueError(f"{series_name} 的基线窗口与末端窗口发生重叠。")
    return baseline_ticks, final_ticks


def compute_agent_shifts(
    series_by_agent: dict[str, dict[int, float]],
    baseline_ticks: list[int],
    final_ticks: list[int],
) -> list[AgentShift]:
    shifts: list[AgentShift] = []
    for agent_id, series in sorted(series_by_agent.items()):
        baseline_values = [series[tick] for tick in baseline_ticks if tick in series]
        final_values = [series[tick] for tick in final_ticks if tick in series]
        if not baseline_values or not final_values:
            continue
        baseline_mean = sum(baseline_values) / len(baseline_values)
        final_mean = sum(final_values) / len(final_values)
        baseline_abs = sum(abs(value) for value in baseline_values) / len(baseline_values)
        final_abs = sum(abs(value) for value in final_values) / len(final_values)
        shifts.append(
            AgentShift(
                agent_id=agent_id,
                baseline_mean_opinion=baseline_mean,
                final_mean_opinion=final_mean,
                delta_opinion=final_mean - baseline_mean,
                baseline_mean_abs_opinion=baseline_abs,
                final_mean_abs_opinion=final_abs,
                delta_abs_opinion=final_abs - baseline_abs,
            )
        )
    return shifts


def summarize_windows(metrics: list[TickMetric], baseline_ticks: list[int], final_ticks: list[int]) -> dict[str, float]:
    by_tick = {metric.tick: metric for metric in metrics}
    baseline_metrics = [by_tick[tick] for tick in baseline_ticks]
    final_metrics = [by_tick[tick] for tick in final_ticks]
    names = [
        "std_population",
        "mean_abs_opinion",
        "pairwise_mean_abs_distance",
        "support_share",
        "oppose_share",
        "neutral_share",
        "extreme_share",
        "two_side_share",
        "polarization_index",
    ]
    out: dict[str, float] = {}
    for name in names:
        baseline_value = metric_mean(baseline_metrics, name)
        final_value = metric_mean(final_metrics, name)
        out[f"baseline_{name}"] = baseline_value
        out[f"final_{name}"] = final_value
        out[f"delta_{name}"] = final_value - baseline_value
    return out


def metric_mean(metrics: Iterable[TickMetric], name: str) -> float:
    values = [float(getattr(metric, name)) for metric in metrics]
    return sum(values) / len(values) if values else 0.0


def one_sided_sign_test(deltas: list[float]) -> float | None:
    positives = sum(1 for delta in deltas if delta > 0)
    negatives = sum(1 for delta in deltas if delta < 0)
    n = positives + negatives
    if n == 0:
        return None
    # 单侧符号检验：原假设下正负变化概率均为 0.5。
    probability = 0.0
    for k in range(positives, n + 1):
        probability += math.comb(n, k) * (0.5 ** n)
    return probability


def bootstrap_mean_ci(deltas: list[float], *, samples: int, seed: int) -> tuple[float | None, float | None]:
    if not deltas or samples == 0:
        return None, None
    rng = random.Random(seed)
    means: list[float] = []
    for _ in range(samples):
        drawn = [deltas[rng.randrange(len(deltas))] for _ in deltas]
        means.append(sum(drawn) / len(drawn))
    means.sort()
    low_index = max(0, int(0.025 * (len(means) - 1)))
    high_index = min(len(means) - 1, int(0.975 * (len(means) - 1)))
    return means[low_index], means[high_index]


def write_metrics_csv(path: Path, metrics: list[TickMetric]) -> None:
    fields = [
        "tick",
        "n_agents",
        "mean_opinion",
        "median_opinion",
        "std_population",
        "mean_abs_opinion",
        "median_abs_opinion",
        "pairwise_mean_abs_distance",
        "min_opinion",
        "max_opinion",
        "range_opinion",
        "support_count",
        "oppose_count",
        "neutral_count",
        "extreme_count",
        "support_share",
        "oppose_share",
        "neutral_share",
        "extreme_share",
        "two_side_share",
        "camp_gap",
        "polarization_index",
    ]
    with path.open("w", newline="", encoding="utf-8") as f:
        writer = csv.DictWriter(f, fieldnames=fields)
        writer.writeheader()
        for metric in metrics:
            writer.writerow({field: format_number(getattr(metric, field)) for field in fields})


def write_agent_shift_csv(path: Path, shifts: list[AgentShift]) -> None:
    fields = [
        "agent_id",
        "baseline_mean_opinion",
        "final_mean_opinion",
        "delta_opinion",
        "baseline_mean_abs_opinion",
        "final_mean_abs_opinion",
        "delta_abs_opinion",
    ]
    with path.open("w", newline="", encoding="utf-8") as f:
        writer = csv.DictWriter(f, fieldnames=fields)
        writer.writeheader()
        for shift in shifts:
            writer.writerow({field: format_number(getattr(shift, field)) for field in fields})


def write_polarization_report(
    *,
    report_path: Path,
    run_dir: Path,
    metrics_path: Path,
    shifts_path: Path,
    options: AnalysisOptions,
    series_by_agent: dict[str, dict[int, float]],
    ticks: list[int],
    baseline_ticks: list[int],
    final_ticks: list[int],
    stats: dict[str, float],
    sign_test_p: float | None,
    ci_low: float | None,
    ci_high: float | None,
) -> None:
    side_pass = stats["final_two_side_share"] >= options.min_side_share
    distance_pass = stats["delta_pairwise_mean_abs_distance"] >= options.min_pairwise_delta
    abs_delta_pass = stats["delta_mean_abs_opinion"] >= options.min_abs_delta
    sign_pass = sign_test_p is not None and sign_test_p <= options.alpha
    ci_pass = ci_low is not None and ci_low > 0.0
    neutral_pass = stats["delta_neutral_share"] <= 0.0
    strict_pass = side_pass and distance_pass and abs_delta_pass and sign_pass and ci_pass and neutral_pass

    lines = [
        "# 舆论极化统计报告",
        "",
        "## 数据范围",
        "",
        f"- 运行目录：`{run_dir}`",
        f"- 智能体数量：{len(series_by_agent)}",
        f"- 分析 tick 数：{len(ticks)}",
        f"- tick 范围：{ticks[0]} - {ticks[-1]}",
        f"- 基线窗口：{baseline_ticks[0]} - {baseline_ticks[-1]}",
        f"- 末端窗口：{final_ticks[0]} - {final_ticks[-1]}",
        "",
        "## 阈值",
        "",
        markdown_table(
            ["项目", "数值"],
            [
                ["明显支持", f"opinion >= {options.support_threshold}"],
                ["明显反对", f"opinion <= {options.oppose_threshold}"],
                ["中立", f"abs(opinion) <= {options.neutral_threshold}"],
                ["极端", f"abs(opinion) >= {options.extreme_threshold}"],
                ["双边阵营最小占比", options.min_side_share],
                ["平均成对距离最小增长", options.min_pairwise_delta],
                ["平均绝对立场最小增长", options.min_abs_delta],
                ["符号检验 alpha", options.alpha],
            ],
        ),
        "",
        "## 前后窗口统计",
        "",
        markdown_table(
            ["指标", "基线窗口均值", "末端窗口均值", "变化"],
            [
                metric_row(stats, "std_population"),
                metric_row(stats, "mean_abs_opinion"),
                metric_row(stats, "pairwise_mean_abs_distance"),
                metric_row(stats, "support_share"),
                metric_row(stats, "oppose_share"),
                metric_row(stats, "neutral_share"),
                metric_row(stats, "extreme_share"),
                metric_row(stats, "two_side_share"),
                metric_row(stats, "polarization_index"),
            ],
        ),
        "",
        "## 配对检验",
        "",
        f"- 单侧符号检验 p 值：{format_number(sign_test_p)}",
        f"- 平均绝对立场变化 bootstrap 95% CI：[{format_number(ci_low)}, {format_number(ci_high)}]",
        "",
        "## 判定标准",
        "",
        markdown_table(
            ["标准", "是否满足"],
            [
                ["末端窗口存在双边阵营", yes_no(side_pass)],
                ["平均成对距离增长达到阈值", yes_no(distance_pass)],
                ["平均绝对立场增长达到阈值", yes_no(abs_delta_pass)],
                ["配对符号检验达到显著性水平", yes_no(sign_pass)],
                ["bootstrap 置信区间下界大于 0", yes_no(ci_pass)],
                ["中立占比没有上升", yes_no(neutral_pass)],
            ],
        ),
        "",
        f"**严格判定：{yes_no(strict_pass)}**",
        "",
        "## 输出文件",
        "",
        f"- `{metrics_path}`",
        f"- `{shifts_path}`",
        "",
    ]
    report_path.write_text("\n".join(lines), encoding="utf-8")


def metric_row(stats: dict[str, float], name: str) -> list[str]:
    return [
        name,
        format_number(stats[f"baseline_{name}"]),
        format_number(stats[f"final_{name}"]),
        format_number(stats[f"delta_{name}"]),
    ]


def plot_history_columns(
    *,
    history_dir: Path,
    columns: list[str],
    output_dir: Path | None,
    x_column: str | None,
    sheet_name: str | None,
    header_row: int,
    image_format: str,
    data_length: int | None,
    recursive: bool,
    show_legend: bool,
    legend_limit: int,
) -> list[Path]:
    table_paths = find_table_files(history_dir, recursive=recursive)
    out_dir = output_dir if output_dir is not None else history_dir / "column_plots"
    series_by_column = collect_series_by_column(
        table_paths=table_paths,
        history_dir=history_dir,
        columns=columns,
        x_column=x_column,
        sheet_name=sheet_name,
        header_row=header_row,
        data_length=data_length,
    )
    return draw_column_plots(
        series_by_column=series_by_column,
        output_dir=out_dir,
        image_format=image_format,
        x_label=x_column or "数据行序号",
        show_legend=show_legend,
        legend_limit=legend_limit,
    )


def find_table_files(history_dir: Path, recursive: bool) -> list[Path]:
    if not history_dir.exists():
        raise FileNotFoundError(f"history 文件夹不存在：{history_dir}")
    if not history_dir.is_dir():
        raise NotADirectoryError(f"history 路径不是文件夹：{history_dir}")

    pattern = "**/*" if recursive else "*"
    table_paths = sorted(
        path
        for path in history_dir.glob(pattern)
        if path.is_file()
        and path.suffix.lower() in TABLE_SUFFIXES
        and path.name not in GENERATED_ANALYSIS_FILES
        and not path.name.startswith("~$")
    )
    if not table_paths:
        raise FileNotFoundError(f"没有在 history 文件夹下找到 CSV/Excel 文件：{history_dir}")
    return table_paths


def print_column_report(table_paths: Iterable[Path], sheet_name: str | None, header_row: int) -> None:
    for table_path in table_paths:
        for current_sheet, headers, _ in iter_sheet_rows(table_path, sheet_name, header_row):
            header_text = ", ".join(format_header(value) for value in headers) or "（无表头）"
            print(f"{table_path} | {current_sheet}")
            print(f"  {header_text}")


def collect_series_by_column(
    *,
    table_paths: Iterable[Path],
    history_dir: Path,
    columns: list[str],
    x_column: str | None,
    sheet_name: str | None,
    header_row: int,
    data_length: int | None,
) -> dict[str, list[PlotSeries]]:
    series_by_column: dict[str, list[PlotSeries]] = {column: [] for column in columns}
    for table_path in table_paths:
        for current_sheet, headers, rows in iter_sheet_rows(table_path, sheet_name, header_row):
            header_index = build_header_index(headers, table_path, current_sheet)
            ensure_required_columns(header_index, headers, table_path, current_sheet, columns, x_column)
            for column in columns:
                y_index = header_index[column]
                x_index = header_index[x_column] if x_column is not None else None
                x_values: list[Any] = []
                y_values: list[float] = []
                for row_offset, row_values in enumerate(rows, start=1):
                    row_number = header_row + row_offset
                    y_raw = cell_value(row_values, y_index)
                    if is_blank(y_raw):
                        continue
                    y_value = parse_number(y_raw, table_path, current_sheet, row_number, column)
                    if x_index is None:
                        x_value = row_offset
                    else:
                        x_value = cell_value(row_values, x_index)
                        if is_blank(x_value):
                            raise ValueError(
                                f"{table_path} | {current_sheet} 第 {row_number} 行横轴列 {x_column!r} 为空。"
                            )
                    x_values.append(x_value)
                    y_values.append(y_value)

                if not y_values:
                    raise ValueError(f"{table_path} | {current_sheet} 列 {column!r} 没有可绘图数据。")
                if data_length is not None:
                    # 绘图范围与统计范围保持一致，保留开头数据。
                    x_values = x_values[:data_length]
                    y_values = y_values[:data_length]
                series_by_column[column].append(
                    PlotSeries(
                        label=series_label(history_dir, table_path, current_sheet),
                        x_values=x_values,
                        y_values=y_values,
                    )
                )

    for column, series_list in series_by_column.items():
        if not series_list:
            raise ValueError(f"列 {column!r} 没有收集到可绘图数据。")
    return series_by_column


def iter_sheet_rows(
    table_path: Path,
    sheet_name: str | None,
    header_row: int,
) -> Iterable[tuple[str, list[Any], list[list[Any]]]]:
    suffix = table_path.suffix.lower()
    if suffix in CSV_SUFFIXES:
        yield from iter_csv_rows(table_path, sheet_name, header_row)
        return
    if suffix in OPENPYXL_SUFFIXES:
        yield from iter_openpyxl_sheet_rows(table_path, sheet_name, header_row)
        return
    if suffix in XLRD_SUFFIXES:
        yield from iter_xlrd_sheet_rows(table_path, sheet_name, header_row)
        return
    raise ValueError(f"不支持的文件后缀：{table_path}")


def iter_csv_rows(
    csv_path: Path,
    sheet_name: str | None,
    header_row: int,
) -> Iterable[tuple[str, list[Any], list[list[Any]]]]:
    if sheet_name is not None:
        raise ValueError(f"{csv_path} 是 CSV 文件，不支持 --sheet。")
    with csv_path.open("r", newline="", encoding="utf-8-sig") as f:
        rows = list(csv.reader(f))
    if len(rows) < header_row:
        yield "CSV", [], []
        return
    # CSV 首行按原始表头精确匹配，不做大小写或空格改写。
    yield "CSV", rows[header_row - 1], rows[header_row:]


def iter_openpyxl_sheet_rows(
    excel_path: Path,
    sheet_name: str | None,
    header_row: int,
) -> Iterable[tuple[str, list[Any], list[list[Any]]]]:
    try:
        from openpyxl import load_workbook
    except ImportError as exc:
        raise RuntimeError("读取 .xlsx/.xlsm 需要安装 openpyxl；请执行 pip install -r requirements.txt。") from exc

    workbook = load_workbook(excel_path, read_only=True, data_only=True)
    try:
        names = workbook.sheetnames
        selected_sheets = selected_sheet_names(names, sheet_name, excel_path)
        for current_sheet in selected_sheets:
            worksheet = workbook[current_sheet]
            header_values = next(
                worksheet.iter_rows(min_row=header_row, max_row=header_row, values_only=True),
                None,
            )
            headers = list(header_values or [])
            rows = [list(row) for row in worksheet.iter_rows(min_row=header_row + 1, values_only=True)]
            yield current_sheet, headers, rows
    finally:
        workbook.close()


def iter_xlrd_sheet_rows(
    excel_path: Path,
    sheet_name: str | None,
    header_row: int,
) -> Iterable[tuple[str, list[Any], list[list[Any]]]]:
    try:
        import xlrd
    except ImportError as exc:
        raise RuntimeError("读取 .xls 需要安装 xlrd；请执行 pip install -r requirements.txt。") from exc

    workbook = xlrd.open_workbook(str(excel_path), on_demand=True)
    selected_sheets = selected_sheet_names(workbook.sheet_names(), sheet_name, excel_path)
    for current_sheet in selected_sheets:
        worksheet = workbook.sheet_by_name(current_sheet)
        if worksheet.nrows < header_row:
            yield current_sheet, [], []
        else:
            yield (
                current_sheet,
                worksheet.row_values(header_row - 1),
                [worksheet.row_values(row_index) for row_index in range(header_row, worksheet.nrows)],
            )


def selected_sheet_names(all_names: list[str], sheet_name: str | None, excel_path: Path) -> list[str]:
    if sheet_name is None:
        return list(all_names)
    if sheet_name not in all_names:
        names = ", ".join(repr(name) for name in all_names)
        raise ValueError(f"{excel_path} 不存在工作表 {sheet_name!r}；实际工作表：{names}")
    return [sheet_name]


def build_header_index(headers: list[Any], table_path: Path, sheet_name: str) -> dict[str, int]:
    header_index: dict[str, int] = {}
    repeated: set[str] = set()
    for index, value in enumerate(headers):
        if not isinstance(value, str) or value == "":
            continue
        if value in header_index:
            repeated.add(value)
            continue
        header_index[value] = index
    if repeated:
        names = ", ".join(repr(value) for value in sorted(repeated))
        raise ValueError(f"{table_path} | {sheet_name} 存在重复表头：{names}。请先改为唯一表头。")
    return header_index


def ensure_required_columns(
    header_index: dict[str, int],
    headers: list[Any],
    table_path: Path,
    sheet_name: str,
    columns: list[str],
    x_column: str | None,
) -> None:
    required = list(columns)
    if x_column is not None:
        required.append(x_column)
    missing = [column for column in required if column not in header_index]
    if missing:
        missing_text = ", ".join(repr(column) for column in missing)
        header_text = ", ".join(format_header(value) for value in headers) or "（无表头）"
        raise ValueError(f"{table_path} | {sheet_name} 缺少列：{missing_text}。实际表头：{header_text}")


def cell_value(row_values: list[Any], index: int) -> Any:
    if index >= len(row_values):
        return None
    return row_values[index]


def parse_number(value: Any, table_path: Path, sheet_name: str, row_number: int, column: str) -> float:
    if isinstance(value, bool):
        raise ValueError(f"{table_path} | {sheet_name} 第 {row_number} 行列 {column!r} 是布尔值，不能绘图。")
    if isinstance(value, (int, float)):
        number = float(value)
    elif isinstance(value, str):
        try:
            number = float(value)
        except ValueError as exc:
            raise ValueError(
                f"{table_path} | {sheet_name} 第 {row_number} 行列 {column!r} 不是数值：{value!r}"
            ) from exc
    else:
        raise ValueError(f"{table_path} | {sheet_name} 第 {row_number} 行列 {column!r} 不是数值：{value!r}")
    if not math.isfinite(number):
        raise ValueError(f"{table_path} | {sheet_name} 第 {row_number} 行列 {column!r} 不是有限数值。")
    return number


def draw_column_plots(
    *,
    series_by_column: dict[str, list[PlotSeries]],
    output_dir: Path,
    image_format: str,
    x_label: str,
    show_legend: bool,
    legend_limit: int,
) -> list[Path]:
    if image_format == "svg":
        return draw_column_svg_plots(
            series_by_column=series_by_column,
            output_dir=output_dir,
            x_label=x_label,
            show_legend=show_legend,
            legend_limit=legend_limit,
        )

    try:
        import matplotlib

        matplotlib.use("Agg")
        import matplotlib.pyplot as plt
    except ImportError as exc:
        raise RuntimeError("绘图需要安装 matplotlib；请执行 pip install -r requirements.txt。") from exc

    plt.rcParams["font.sans-serif"] = ["Microsoft YaHei", "SimHei", "Arial Unicode MS", "DejaVu Sans"]
    plt.rcParams["axes.unicode_minus"] = False
    output_dir.mkdir(parents=True, exist_ok=True)
    used_names: set[str] = set()
    output_paths: list[Path] = []
    for column, series_list in series_by_column.items():
        fig, ax = plt.subplots(figsize=(11, 6.5))
        for series in series_list:
            ax.plot(series.x_values, series.y_values, marker="o", linewidth=1.7, markersize=3, label=series.label)
        ax.set_title(column)
        ax.set_xlabel(x_label)
        ax.set_ylabel(column)
        ax.grid(True, linestyle="--", linewidth=0.5, alpha=0.4)
        if show_legend and len(series_list) <= legend_limit:
            ax.legend(loc="center left", bbox_to_anchor=(1.02, 0.5), fontsize=8)
        fig.tight_layout()
        output_path = output_dir / unique_output_name(column, image_format, used_names)
        fig.savefig(output_path, dpi=150)
        plt.close(fig)
        output_paths.append(output_path)
    return output_paths


def draw_column_svg_plots(
    *,
    series_by_column: dict[str, list[PlotSeries]],
    output_dir: Path,
    x_label: str,
    show_legend: bool,
    legend_limit: int,
) -> list[Path]:
    output_dir.mkdir(parents=True, exist_ok=True)
    used_names: set[str] = set()
    output_paths: list[Path] = []
    for column, series_list in series_by_column.items():
        output_path = output_dir / unique_output_name(column, "svg", used_names)
        write_column_svg_plot(
            output_path=output_path,
            title=column,
            x_label=x_label,
            y_label=column,
            series_list=series_list,
            show_legend=show_legend and len(series_list) <= legend_limit,
        )
        output_paths.append(output_path)
    return output_paths


def write_column_svg_plot(
    *,
    output_path: Path,
    title: str,
    x_label: str,
    y_label: str,
    series_list: list[PlotSeries],
    show_legend: bool,
) -> None:
    width = 1120
    height = 660
    left = 82
    right = 290 if show_legend else 36
    top = 62
    bottom = 104
    plot_width = width - left - right
    plot_height = height - top - bottom
    colors = ("#2563eb", "#dc2626", "#16a34a", "#9333ea", "#ea580c", "#0891b2", "#be123c", "#4d7c0f")
    x_axis = build_x_axis(series_list)
    y_min, y_max = value_bounds([value for series in series_list for value in series.y_values])

    def x_pos(value: Any, fallback_index: int) -> float:
        if x_axis["kind"] == "numeric":
            numeric_value = axis_number(value)
            if numeric_value is None:
                numeric_value = float(fallback_index)
            x_min = x_axis["min"]
            x_max = x_axis["max"]
            return scale_value(numeric_value, x_min, x_max, left, left + plot_width)
        key = display_value(value)
        index = x_axis["index"].get(key, fallback_index)
        count = max(1, len(x_axis["labels"]) - 1)
        return scale_value(index, 0, count, left, left + plot_width)

    def y_pos(value: float) -> float:
        return scale_value(value, y_min, y_max, top + plot_height, top)

    lines = [
        '<?xml version="1.0" encoding="UTF-8"?>',
        f'<svg xmlns="http://www.w3.org/2000/svg" width="{width}" height="{height}" viewBox="0 0 {width} {height}">',
        "<style>text{font-family:Arial,'Microsoft YaHei','SimHei',sans-serif;fill:#111827}.title{font-size:22px;font-weight:700}.label{font-size:14px}.tick{font-size:12px;fill:#4b5563}.axis{stroke:#374151;stroke-width:1.2}.grid{stroke:#d1d5db;stroke-width:0.8;stroke-dasharray:4 4}.legend{font-size:12px}</style>",
        f'<rect x="0" y="0" width="{width}" height="{height}" fill="#ffffff"/>',
        f'<text class="title" x="{width / 2}" y="36" text-anchor="middle">{escape_text(title)}</text>',
    ]
    for tick in numeric_ticks(y_min, y_max, 6):
        y = y_pos(tick)
        lines.append(f'<line class="grid" x1="{left}" y1="{y:.2f}" x2="{left + plot_width}" y2="{y:.2f}"/>')
        lines.append(f'<text class="tick" x="{left - 12}" y="{y + 4:.2f}" text-anchor="end">{format_axis_number(tick)}</text>')
    for axis_tick in x_axis_ticks(x_axis):
        x = x_pos(axis_tick["value"], axis_tick["index"])
        label = escape_text(short_text(axis_tick["label"], 18))
        lines.append(f'<line class="grid" x1="{x:.2f}" y1="{top}" x2="{x:.2f}" y2="{top + plot_height}"/>')
        lines.append(
            f'<text class="tick" x="{x:.2f}" y="{top + plot_height + 30}" text-anchor="end" '
            f'transform="rotate(-28 {x:.2f} {top + plot_height + 30})">{label}</text>'
        )
    lines.extend([
        f'<line class="axis" x1="{left}" y1="{top}" x2="{left}" y2="{top + plot_height}"/>',
        f'<line class="axis" x1="{left}" y1="{top + plot_height}" x2="{left + plot_width}" y2="{top + plot_height}"/>',
        f'<text class="label" x="{left + plot_width / 2}" y="{height - 28}" text-anchor="middle">{escape_text(x_label)}</text>',
        f'<text class="label" x="24" y="{top + plot_height / 2}" text-anchor="middle" transform="rotate(-90 24 {top + plot_height / 2})">{escape_text(y_label)}</text>',
    ])
    for index, series in enumerate(series_list):
        color = colors[index % len(colors)]
        points = " ".join(
            f"{x_pos(x_value, point_index):.2f},{y_pos(y_value):.2f}"
            for point_index, (x_value, y_value) in enumerate(zip(series.x_values, series.y_values))
        )
        if points:
            lines.append(f'<polyline points="{points}" fill="none" stroke="{color}" stroke-width="2.1" stroke-linejoin="round" stroke-linecap="round"/>')
        for point_index, (x_value, y_value) in enumerate(zip(series.x_values, series.y_values)):
            lines.append(f'<circle cx="{x_pos(x_value, point_index):.2f}" cy="{y_pos(y_value):.2f}" r="2.8" fill="{color}"/>')
        if show_legend:
            legend_x = left + plot_width + 28
            legend_y = top + 18 + index * 20
            if legend_y < height - 24:
                lines.append(f'<line x1="{legend_x}" y1="{legend_y}" x2="{legend_x + 22}" y2="{legend_y}" stroke="{color}" stroke-width="2.1"/>')
                lines.append(f'<text class="legend" x="{legend_x + 30}" y="{legend_y + 4}">{escape_text(short_text(series.label, 34))}</text>')
    lines.append("</svg>")
    output_path.write_text("\n".join(lines), encoding="utf-8")


def build_x_axis(series_list: list[PlotSeries]) -> dict[str, Any]:
    all_values = [value for series in series_list for value in series.x_values]
    numeric_values = [axis_number(value) for value in all_values]
    if all(value is not None for value in numeric_values):
        x_min, x_max = value_bounds([value for value in numeric_values if value is not None])
        return {"kind": "numeric", "min": x_min, "max": x_max}
    labels: list[str] = []
    seen: set[str] = set()
    for value in all_values:
        label = display_value(value)
        if label in seen:
            continue
        seen.add(label)
        labels.append(label)
    return {"kind": "category", "labels": labels, "index": {label: index for index, label in enumerate(labels)}}


def x_axis_ticks(x_axis: dict[str, Any]) -> list[dict[str, Any]]:
    if x_axis["kind"] == "numeric":
        ticks = numeric_ticks(x_axis["min"], x_axis["max"], 8)
        return [{"value": tick, "label": format_axis_number(tick), "index": index} for index, tick in enumerate(ticks)]
    labels = x_axis["labels"]
    if not labels:
        return []
    if len(labels) <= 8:
        indexes = list(range(len(labels)))
    else:
        indexes = sorted({round(index * (len(labels) - 1) / 7) for index in range(8)})
    return [{"value": labels[index], "label": labels[index], "index": index} for index in indexes]


def value_bounds(values: list[float]) -> tuple[float, float]:
    low = min(values)
    high = max(values)
    if low == high:
        padding = 1.0 if low == 0 else abs(low) * 0.1
        return low - padding, high + padding
    padding = (high - low) * 0.06
    return low - padding, high + padding


def clamp_number(value: float, low: float, high: float) -> float:
    """把数值限制在给定闭区间内。"""

    return min(high, max(low, value))


def scale_value(value: float, in_min: float, in_max: float, out_min: float, out_max: float) -> float:
    """线性映射坐标；输入范围退化时返回输出中点。"""

    if in_min == in_max:
        return (out_min + out_max) / 2
    ratio = (value - in_min) / (in_max - in_min)
    return out_min + ratio * (out_max - out_min)


def numeric_ticks(low: float, high: float, count: int) -> list[float]:
    if count <= 1 or low == high:
        return [low]
    step = (high - low) / (count - 1)
    return [low + step * index for index in range(count)]


def axis_number(value: Any) -> float | None:
    if isinstance(value, bool):
        return None
    if isinstance(value, (int, float)):
        number = float(value)
        return number if math.isfinite(number) else None
    return None


def display_value(value: Any) -> str:
    return "" if value is None else str(value)


def series_label(history_dir: Path, table_path: Path, sheet_name: str) -> str:
    try:
        relative_path = table_path.relative_to(history_dir)
    except ValueError:
        relative_path = table_path
    return f"{relative_path} | {sheet_name}"


def cell_text(value: object) -> str:
    return "" if value is None else str(value)


def parse_float(value: object) -> float:
    try:
        return float(cell_text(value))
    except (TypeError, ValueError):
        return 0.0


def json_obj(value: object) -> dict:
    if isinstance(value, dict):
        return value
    if not value:
        return {}
    try:
        parsed = json.loads(str(value))
    except (TypeError, json.JSONDecodeError):
        return {}
    return parsed if isinstance(parsed, dict) else {}


def is_blank(value: Any) -> bool:
    return value is None or value == ""


def format_header(value: Any) -> str:
    if value is None or value == "":
        return "（空）"
    return repr(value)


def markdown_table(headers: list[str], rows: list[list[object]]) -> str:
    if not rows:
        rows = [["" for _ in headers]]
    out = [
        "| " + " | ".join(headers) + " |",
        "| " + " | ".join("---" for _ in headers) + " |",
    ]
    for row in rows:
        out.append("| " + " | ".join(str(value) for value in row) + " |")
    return "\n".join(out)


def yes_no(value: bool) -> str:
    return "是" if value else "否"


def format_number(value: object) -> str:
    if value is None:
        return ""
    if isinstance(value, str):
        return value
    if isinstance(value, int):
        return str(value)
    try:
        number = float(value)
    except (TypeError, ValueError):
        return str(value)
    if not math.isfinite(number):
        return ""
    return f"{number:.6f}".rstrip("0").rstrip(".")


def format_axis_number(value: float) -> str:
    if value == 0:
        return "0"
    if abs(value) >= 1000 or abs(value) < 0.01:
        return f"{value:.2e}"
    return f"{value:.3f}".rstrip("0").rstrip(".")


def short_text(value: str, max_length: int) -> str:
    if len(value) <= max_length:
        return value
    return value[: max(0, max_length - 3)] + "..."


def escape_text(value: str) -> str:
    return html.escape(value, quote=True)


def unique_output_name(column: str, image_format: str, used_names: set[str]) -> str:
    base = safe_filename(column) or "plot"
    name = f"{base}.{image_format}"
    index = 2
    while name.lower() in used_names:
        name = f"{base}_{index}.{image_format}"
        index += 1
    used_names.add(name.lower())
    return name


def safe_filename(value: str) -> str:
    # 文件名只替换 Windows 禁止字符，不改动列名匹配逻辑。
    text = "".join("_" if char in INVALID_FILENAME_CHARS else char for char in value)
    return text.strip(" .")


if __name__ == "__main__":
    try:
        raise SystemExit(main())
    except Exception as exc:
        print(f"错误：{exc}", file=sys.stderr)
        raise SystemExit(1)
